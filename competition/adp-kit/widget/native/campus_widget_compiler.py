"""Compiler for real 02/03/04 workflows and six exported native Widgets."""

import copy
import hashlib
import io
import json
import uuid
from pathlib import Path


ROOT = Path(__file__).resolve().parent
CATALOG = json.loads((ROOT / "real-adp-export-catalog.json").read_text(encoding="utf-8"))
REGISTRY = json.loads((ROOT / "widget-registry.json").read_text(encoding="utf-8"))
FIXED_NAMESPACE = uuid.UUID("9b65e3e9-b60b-4c4c-b53d-26f2c847cf4e")

SPECS = {
    "02": {
        "kind": "Classroom", "file": "02-Classroom-Final.zip", "name": "02-空教室规划-Final",
        "tool": "空教室查询", "verify": "结果核验与呈现", "answer": "空教室结果回复",
        "adapter": "classroom-widget-adapter.py", "examples": ["校区A第1周周一第5-6节有哪些空教室"],
    },
    "03": {
        "kind": "Conflict", "file": "03-Conflict-Final.zip", "name": "03-课程冲突比较-Final",
        "tool": "课程冲突比较", "verify": "冲突结果核验与呈现", "answer": "冲突比较结果回复",
        "adapter": "conflict-widget-adapter.py", "examples": [
            "检查教师003第1周周一是否存在时间冲突或跨校区赶场",
            "教师003和教师003第1周周一是否存在冲突或跨校区赶场",
        ],
    },
    "04": {
        "kind": "DayPlan", "file": "04-DayPlan-Final.zip", "name": "04-今日校园计划-Final",
        "tool": "今日计划", "verify": "今日计划核验与呈现", "answer": "今日计划结果回复",
        "adapter": "dayplan-widget-adapter.py", "examples": ["帮我看看2026-09-04的安排"],
    },
}


def node_id(key, role):
    return str(uuid.uuid5(FIXED_NAMESPACE, f"campus-widget-v1:{key}:{role}"))


def type_name(schema):
    value = schema.get("type")
    if value == "integer": return "INT"
    if value == "number": return "FLOAT"
    if value == "boolean": return "BOOL"
    if value == "object": return "OBJECT"
    if value == "array":
        item_type = schema.get("items", {}).get("type")
        return {"string": "ARRAY_STRING", "integer": "ARRAY_INT", "number": "ARRAY_FLOAT"}.get(item_type, "ARRAY_OBJECT")
    return "STRING"


def output_property(name, schema):
    properties = [output_property(key, child) for key, child in schema.get("properties", {}).items()]
    if schema.get("type") == "array" and schema.get("items", {}).get("type") == "object":
        properties = [output_property(key, child) for key, child in schema["items"].get("properties", {}).items()]
    return {"Title": name, "Type": type_name(schema), "Required": [], "Properties": properties,
            "Desc": "", "AnalysisMethod": "COVER"}


def node_ui(node, position):
    inputs = [item.get("Name") for item in node.get("Inputs", [])]
    output_names = ["Output"]
    if node.get("Outputs"):
        output_names.extend(f"Output.{item['Title']}" for item in node["Outputs"][0].get("Properties", []))
    return json.dumps({
        "data": {"isHovering": False, "isParallel": False, "source": True, "target": True,
                 "debug": None, "error": False, "output": [], "schema": None, "checkDataError": 0,
                 "showTips": False, "isConcurrent": False, "financeType": None,
                 "content": {"inputs": inputs, "outputs": output_names}},
        "position": position, "targetPosition": "left", "sourcePosition": "right", "selected": False,
        "measured": {"width": 250, "height": 124}, "dragging": False,
    }, ensure_ascii=False, separators=(",", ":"))


def reference_input(name, source):
    return {"Name": name, "Type": "OBJECT", "Input": {"InputType": "REFERENCE_OUTPUT", "Reference": {
        "NodeID": source, "JsonPath": "Output.Body"}}, "Desc": "verified CampusTools output",
        "IsRequired": False, "SubInputs": [], "DefaultValue": "", "DefaultFileName": ""}


def code_node(key, role, name, source, schema, tool_id, x, code_file):
    nid = node_id(key, role)
    properties = [output_property("route", {"type": "string"})]
    properties.extend(output_property(field, child) for field, child in schema["properties"].items())
    node = {
        "NodeID": nid, "NodeName": name, "NodeDesc": "Deterministic adapter; maps verified Tool facts only.",
        "NodeType": "CODE_EXECUTOR", "CodeExecutorNodeData": {
            "Code": (ROOT / code_file).read_text(encoding="utf-8"), "Language": "PYTHON3"},
        "Inputs": [reference_input("tool_body", tool_id)],
        "Outputs": [{"Title": "Output", "Type": "OBJECT", "Required": [], "Properties": properties,
                     "Desc": "native Widget view model", "AnalysisMethod": "COVER"}],
        "NextNodeIDs": [], "ExceptionHandling": {"Switch": "OFF", "MaxRetries": "1", "RetryInterval": "1",
            "AbnormalOutputResult": "", "HandleMethod": "EXCEPTION_OUTPUT", "NextNodeIDs": [],
            "AbnormalRetrySwitch": "ABNORMAL_RETRY_ON", "Timeout": "10"},
    }
    node["NodeUI"] = node_ui(node, {"x": x, "y": 260 if role == "primary-adapter" else 420})
    return node


def literal_param(name, schema, required=False):
    param_type = type_name(schema)
    subparams = []
    if schema.get("type") == "object":
        child_required = schema.get("required", [])
        subparams = [literal_param(key, child, key in child_required) for key, child in schema.get("properties", {}).items()]
    elif schema.get("type") == "array" and schema.get("items", {}).get("type") == "object":
        item = schema["items"]
        subparams = [{
            "ParamName": "item 0", "ParamDesc": "exported array item schema", "ParamType": "OBJECT",
            "Input": {"InputType": "USER_INPUT", "UserInputValue": {"Values": [""], "FileNames": []}},
            "IsRequired": False,
            "SubParams": [literal_param(key, child, key in item.get("required", [])) for key, child in item.get("properties", {}).items()],
            "IsSecretField": False, "ParamTypeOf": [],
        }]
    return {"ParamName": name, "ParamDesc": "exported Widget schema field", "ParamType": param_type,
        "Input": {"InputType": "USER_INPUT", "UserInputValue": {"Values": [""], "FileNames": []}},
        "IsRequired": required, "SubParams": subparams, "IsSecretField": False, "ParamTypeOf": []}


def widget_param(name, schema, source):
    param = literal_param(name, schema, name in schema.get("_required", []))
    param["Input"] = {"InputType": "REFERENCE_OUTPUT", "Reference": {"NodeID": source, "JsonPath": f"Output.{name}"}}
    return param


def widget_node(key, role, kind, schema, source, x):
    record = REGISTRY["widgets"][kind]
    required = schema.get("required", [])
    properties = copy.deepcopy(schema["properties"])
    for child in properties.values(): child["_required"] = required
    nid = node_id(key, role)
    params = [widget_param(name, child, source) for name, child in properties.items()]
    node = {
        "NodeID": nid, "NodeName": record["name"], "NodeDesc": "Inputs derived from the real exported Widget Schema.",
        "NodeType": "WIDGET", "WidgetNodeData": {"WidgetID": record["widgetId"],
            "ActionType": "WIDGET_ACTION_NONE", "WidgetParam": params},
        "Inputs": [], "Outputs": [], "NextNodeIDs": [],
    }
    node["NodeUI"] = json.dumps({"data": {"content": {"inputs": list(properties), "outputs": ["Output", "Output.Content"]},
        "isHovering": False, "isParallel": False, "source": True, "target": True, "debug": None, "error": False,
        "output": [], "schema": None, "checkDataError": 0, "showTips": False, "isConcurrent": False, "financeType": None},
        "position": {"x": x, "y": 260}, "targetPosition": "left", "sourcePosition": "right", "selected": False,
        "measured": {"width": 250, "height": 124}, "dragging": False}, ensure_ascii=False, separators=(",", ":"))
    return node


def set_literal_value(param, value):
    param["Input"] = {"InputType": "USER_INPUT", "UserInputValue": {
        "Values": [value], "FileNames": []}}


def static_missing_param_widget(key, existing, schema):
    """Replace a pre-tool clarification answer with the real Error Widget."""
    record = REGISTRY["widgets"]["Error"]
    required = schema.get("required", [])
    params = [literal_param(name, child, name in required)
              for name, child in schema["properties"].items()]
    by_name = {param["ParamName"]: param for param in params}
    for name, value in {
        "title": "需要补充比较对象",
        "subtitle": "请补充另一个班级、教师或教室后继续检查。",
        "queryId": "",
        "dataVersion": "competition-demo-v1",
    }.items():
        set_literal_value(by_name[name], value)
    error_fields = {param["ParamName"]: param for param in by_name["error"]["SubParams"]}
    set_literal_value(error_fields["code"], "MISSING_PARAM")
    set_literal_value(error_fields["message"], "课程冲突比较需要两个明确的比较对象。")
    action_fields = {param["ParamName"]: param
                     for param in by_name["actions"]["SubParams"][0]["SubParams"]}
    for name, value in {
        "id": "supply-second-object",
        "type": "sys.chat",
        "label": "补充比较对象",
        "message": "请补充另一个需要比较的班级、教师或教室",
    }.items():
        set_literal_value(action_fields[name], value)
    node = {
        "NodeID": existing["NodeID"],
        "NodeName": f"{record['name']}-缺少比较对象",
        "NodeDesc": "Pre-tool MISSING_PARAM rendered by the real exported Error Widget.",
        "NodeType": "WIDGET",
        "WidgetNodeData": {
            "WidgetID": record["widgetId"],
            "ActionType": "WIDGET_ACTION_NONE",
            "WidgetParam": params,
        },
        "Inputs": [], "Outputs": [], "NextNodeIDs": existing.get("NextNodeIDs", []),
    }
    position = json.loads(existing.get("NodeUI") or "{}").get("position", {"x": 800, "y": 80})
    node["NodeUI"] = json.dumps({
        "data": {"content": {"inputs": list(schema["properties"]), "outputs": ["Output", "Output.Content"]},
                 "isHovering": False, "isParallel": False, "source": True, "target": True,
                 "debug": None, "error": False, "output": [], "schema": None, "checkDataError": 0,
                 "showTips": False, "isConcurrent": False, "financeType": None},
        "position": position, "targetPosition": "left", "sourcePosition": "right", "selected": False,
        "measured": {"width": 250, "height": 124}, "dragging": False,
    }, ensure_ascii=False, separators=(",", ":"))
    return node


def logic_node(key, adapter_id, recovery_id, targets, x):
    nid = node_id(key, "widget-router")
    routes = [("primary", targets["primary"]), ("choice", targets["choice"]), ("error", targets["error"])]
    groups, ui = [], []
    handles = {}
    for index, (route, target) in enumerate(routes):
        handle = node_id(key, f"route-{route}")
        handles[route] = handle
        groups.append({"NextNodeIDs": [target], "Logical": {"LogicalOperator": "UNSPECIFIED", "Compound": [],
            "Comparison": {"Left": {"InputType": "REFERENCE_OUTPUT", "Reference": {"NodeID": recovery_id, "JsonPath": "Output.route"}},
                "LeftType": "STRING", "Operator": "EQ", "Right": {"InputType": "USER_INPUT", "UserInputValue": {"Values": [route], "FileNames": []}}, "MatchType": "SEMANTIC"}}})
        ui.append({"content": [{"leftStr": "Output.route", "rightStr": route, "operatorStr": "等于"}], "index": index, "id": handle})
    fallback = node_id(key, "route-fallback")
    handles["fallback"] = fallback
    groups.append({"NextNodeIDs": [targets["fallback"]]})
    ui.append({"content": [], "index": 3, "id": fallback})
    node = {"NodeID": nid, "NodeName": "Native Widget Route", "NodeDesc": "primary/choice/error/fallback deterministic route",
        "NodeType": "LOGIC_EVALUATOR", "LogicEvaluatorNodeData": {"Group": groups, "ModelParams": {"Temperature": 0}},
        "Inputs": [], "Outputs": [{"Title": "Output", "Type": "OBJECT", "Required": [], "Properties": [{
            "Title": "ConditionIndex", "Type": "INT", "Required": [], "Properties": [], "Desc": "selected route", "AnalysisMethod": "COVER"}],
            "Desc": "route", "AnalysisMethod": "COVER"}], "NextNodeIDs": []}
    node["NodeUI"] = json.dumps({"data": {"content": ui, "isHovering": False, "isParallel": False, "source": True, "target": True,
        "debug": None, "error": False, "output": [], "schema": None, "checkDataError": 0, "showTips": False, "isConcurrent": False,
        "financeType": None}, "position": {"x": x, "y": 340}, "targetPosition": "left", "sourcePosition": "right",
        "selected": False, "measured": {"width": 250, "height": 124}, "dragging": False}, ensure_ascii=False, separators=(",", ":"))
    return node, handles


def edge(source, target, handle=None):
    source_handle = handle or f"{source}-source"
    return {"source": source, "target": target, "sourceHandle": source_handle, "targetHandle": f"{target}-target", "type": "custom",
        "data": {"connectedNodeIsHovering": False, "error": False, "isHovering": False},
        "id": f"xy-edge__{source}{source_handle}-{target}{target}-target", "selected": False, "animated": False}


def compile_workflow(key, workflow):
    spec = SPECS[key]
    nodes = workflow["Nodes"]
    by_name = {node["NodeName"]: node for node in nodes}
    tool, verify, answer = by_name[spec["tool"]], by_name[spec["verify"]], by_name[spec["answer"]]
    end = next(node for node in nodes if node["NodeType"] == "END")
    workflow["WorkflowName"] = spec["name"]
    workflow["WorkflowDesc"] = f"Campus Widget Compiler: verified {spec['kind']} facts → deterministic Adapter → real native Widget; Choice/Error recovery enabled."
    workflow["ReleaseTime"] = ""
    widget_schema = json.loads((ROOT / "contracts" / spec["kind"].lower() / "schema.json").read_text(encoding="utf-8"))
    choice_schema = json.loads((ROOT / "contracts" / "choice" / "schema.json").read_text(encoding="utf-8"))
    error_schema = json.loads((ROOT / "contracts" / "error" / "schema.json").read_text(encoding="utf-8"))
    if key == "03" and "缺少比较对象" in by_name:
        missing = by_name["缺少比较对象"]
        nodes[nodes.index(missing)] = static_missing_param_widget(key, missing, error_schema)
    base_x = max(json.loads(node["NodeUI"])["position"]["x"] for node in nodes if node.get("NodeUI")) + 260
    primary = code_node(key, "primary-adapter", f"{spec['kind']} Adapter", spec["kind"], widget_schema, tool["NodeID"], base_x, spec["adapter"])
    recovery_schema = copy.deepcopy(choice_schema)
    recovery_schema["properties"]["error"] = copy.deepcopy(error_schema["properties"]["error"])
    recovery = code_node(key, "recovery-adapter", "Choice/Error Adapter", "Recovery", recovery_schema, tool["NodeID"], base_x + 280, "recovery-widget-adapter.py")
    primary_widget = widget_node(key, "primary-widget", spec["kind"], widget_schema, primary["NodeID"], base_x + 840)
    choice_widget = widget_node(key, "choice-widget", "Choice", choice_schema, recovery["NodeID"], base_x + 840)
    error_widget = widget_node(key, "error-widget", "Error", error_schema, recovery["NodeID"], base_x + 840)
    router, handles = logic_node(key, primary["NodeID"], recovery["NodeID"], {
        "primary": primary_widget["NodeID"], "choice": choice_widget["NodeID"], "error": error_widget["NodeID"], "fallback": answer["NodeID"],
    }, base_x + 560)
    verify["NextNodeIDs"] = [primary["NodeID"]]
    primary["NextNodeIDs"] = [recovery["NodeID"]]
    recovery["NextNodeIDs"] = [router["NodeID"]]
    for widget in (primary_widget, choice_widget, error_widget): widget["NextNodeIDs"] = [end["NodeID"]]
    nodes.extend([primary, recovery, router, primary_widget, choice_widget, error_widget])
    edges = json.loads(workflow["Edge"]) if isinstance(workflow["Edge"], str) else workflow["Edge"]
    edges = [item for item in edges if not (item.get("source") == verify["NodeID"] and item.get("target") == answer["NodeID"])]
    edges.extend([edge(verify["NodeID"], primary["NodeID"]), edge(primary["NodeID"], recovery["NodeID"]), edge(recovery["NodeID"], router["NodeID"])])
    for route, target in [("primary", primary_widget), ("choice", choice_widget), ("error", error_widget)]:
        edges.append(edge(router["NodeID"], target["NodeID"], f"{router['NodeID']}.{handles[route]}-source"))
        edges.append(edge(target["NodeID"], end["NodeID"]))
    edges.append(edge(router["NodeID"], answer["NodeID"], f"{router['NodeID']}.{handles['fallback']}-source"))
    workflow["Edge"] = json.dumps(edges, ensure_ascii=False, separators=(",", ":"))
    return workflow


def source_path(key, downloads):
    record = CATALOG["workflows"][key]
    direct = downloads / record["file"]
    for path in [direct, *downloads.glob(f"export-{key}-*.zip")]:
        if path.is_file() and hashlib.sha256(path.read_bytes()).hexdigest() == record["sha256"]:
            return path
    raise AssertionError(f"real workflow export missing or hash drift: {key}")


def workbook_rows_from_zip(seed_path, name):
    from openpyxl import load_workbook
    import zipfile
    with zipfile.ZipFile(seed_path) as reader:
        workbook = load_workbook(io.BytesIO(reader.read(name)), read_only=True, data_only=False)
    rows = list(workbook.active.iter_rows(values_only=True))
    return list(rows[0]), [list(row) for row in rows[1:]]


def workbook_records(seed_path, workflow, examples, deterministic_xlsx):
    workflow_id = workflow["WorkflowID"]
    payloads = {}
    for name in ["parameters.xlsx", "variables.xlsx", "workflow_references.xlsx"]:
        headers, rows = workbook_rows_from_zip(seed_path, name)
        if "WorkflowId" in headers:
            index = headers.index("WorkflowId")
            for row in rows: row[index] = workflow_id
        payloads[name] = deterministic_xlsx(headers, rows)
    payloads["workflows.xlsx"] = deterministic_xlsx(
        ["WorkflowId", "WorkflowName", "WorkflowDescription", "CanvasStructure"],
        [[workflow_id, workflow["WorkflowName"], workflow["WorkflowDesc"], f"{workflow_id}_workflow.json"]],
    )
    payloads["example_queries.xlsx"] = deterministic_xlsx(
        ["WorkflowId", "ExampleQueryContent"], [[workflow_id, example] for example in examples],
    )
    return payloads
