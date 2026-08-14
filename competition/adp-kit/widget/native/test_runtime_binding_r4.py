import io
import json
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[3]
OUTPUT = REPO / "output" / "competition-adp" / "r4"
sys.path.insert(0, str(ROOT))

from bind_runtime_environment import (  # noqa: E402
    BASE_WORKFLOW_SLOTS,
    WORKFLOW_FILES,
    WORKFLOW_NAMES,
    assert_current_runtime_id,
    workflow_from_zip,
)


EXPECTED_IDS = {
    "Schedule": "23fbc659efe3482fab588d754e4420a4",
    "Classroom": "d781773ab49c4b15a5a4b99f0e28a748",
    "Conflict": "86e9ea7eb9e0458cb69a5d08aa767163",
    "DayPlan": "9e334031b8bc47e9ab349244f9bca421",
    "Choice": "b2726b1796294285b9f2daffc6d75dc7",
    "Error": "906df4c8d24548fabb8bd8a080de76d2",
}
OLD_IDS = {
    "9eb1ec5e1deb416ab5aba320359439a5",
    "8d576e5af9b04fdd99804e7fcbff3644",
    "f4ff76029d8142caa86fb9baac5314bf",
    "f540588933a4459cbe78a6fe99aa022c",
    "3161078fd4d54a27bae65f38cd44c537",
}
HEADERS = {
    "workflows.xlsx": ["WorkflowId", "WorkflowName", "WorkflowDescription", "CanvasStructure"],
    "parameters.xlsx": ["WorkflowId", "WorkflowNodeId", "WorkflowNodeName", "ParameterId", "ParameterName", "ParameterDescription", "ParameterType", "ParameterCorrectExample", "ParameterIncorrectExample", "ParameterParentId", "ParameterRequiredStatus"],
    "example_queries.xlsx": ["WorkflowId", "ExampleQueryContent"],
    "variables.xlsx": ["WorkflowId", "VariableId", "VariableName", "VariableDescription", "VariableType", "VariableDefaultValue", "VariableDefaultValueFileName", "ParameterType"],
    "workflow_references.xlsx": ["WorkflowId", "WorkflowNodeId", "ReferenceWorkflowId", "ReferenceWorkflowName"],
}


def expect_red(callable_value, expected_text):
    try:
        callable_value()
    except AssertionError as error:
        assert expected_text in str(error), error
        return
    raise AssertionError(f"expected RED gate: {expected_text}")


def run_wrapped_adapter(code, key):
    namespace = {}
    exec(code, namespace)
    bodies = {
        "02": {
            "success": True, "dataVersion": "competition-demo-v1", "queryId": "q-r4",
            "evidence": {"verified": True}, "query": {"week": 1, "weekday": 1, "periodStart": 5, "periodEnd": 6, "campus": "校区A"},
            "items": [],
        },
        "03": {
            "success": True, "dataVersion": "competition-demo-v1", "queryId": "q-r4",
            "evidence": {"verified": True}, "query": {"week": 1, "weekday": 1},
            "items": [], "rushWarnings": [], "summary": {"selfCompare": True, "conflictCount": 0, "rushWarningCount": 0},
            "compared": [{"type": "teacher", "name": "教师003"}, {"type": "teacher", "name": "教师003"}],
        },
        "04": {
            "success": True, "dataVersion": "competition-demo-v1", "queryId": "q-r4",
            "evidence": {"verified": True}, "query": {"date": "2026-09-04", "week": 1, "weekday": 5},
            "items": [], "summary": {"lessonCount": 0}, "resolvedEntity": {"type": "user", "name": "匿名演示用户"},
        },
    }
    output = namespace["main"]({"tool_body": bodies[key]})
    messages = [value for name, value in output.items() if name.endswith("Message") and isinstance(value, str)]
    assert messages, "wrapped adapter must emit sys.chat messages"
    for message in messages:
        assert "【小序ActionV2】" in message, message
        assert "query=" in message and "intent=" in message, message
    return output


def main():
    registry = json.loads((OUTPUT / "widget-registry.runtime.json").read_text(encoding="utf-8"))
    assert registry["environment"] == "competition-current"
    assert registry["bindingPolicy"] == "FAIL_CLOSED_CURRENT_ENVIRONMENT_EXPORTS_ONLY"
    assert {key: value["realWidgetId"] for key, value in registry["widgets"].items()} == EXPECTED_IDS
    assert all(value["exportSha256"] and value["viewSha256"] and value["schemaSha256"] for value in registry["widgets"].values())
    assert all(value["exportedAt"] for value in registry["widgets"].values())

    expect_red(lambda: assert_current_runtime_id("empty.widget", "", OLD_IDS), "empty or invalid")
    expect_red(
        lambda: assert_current_runtime_id("stale.widget", next(iter(OLD_IDS)), OLD_IDS),
        "stale logical-registry WidgetID",
    )

    allowed_ids = set(EXPECTED_IDS.values())
    workflow_ids = set()
    for key in BASE_WORKFLOW_SLOTS:
        file_name = WORKFLOW_FILES[key]
        artifact = OUTPUT / file_name
        workflow, _ = workflow_from_zip(artifact)
        assert workflow["WorkflowName"] == WORKFLOW_NAMES[key]
        assert workflow["WorkflowID"] not in workflow_ids
        workflow_ids.add(workflow["WorkflowID"])
        widgets = [node for node in workflow["Nodes"] if node.get("NodeType") == "WIDGET"]
        assert widgets
        for node in widgets:
            data = node["WidgetNodeData"]
            assert data["WidgetID"] in allowed_ids
            assert data["WidgetID"] != ""
            assert all(param.get("ParamType") in {"STRING", "INT", "FLOAT", "BOOL"} for param in data.get("WidgetParam", []))
            if data["WidgetID"] == EXPECTED_IDS["Error"]:
                assert "queryId" not in [param.get("ParamName") for param in data.get("WidgetParam", [])]
            for param in data.get("WidgetParam", []):
                if not str(param.get("ParamName") or "").endswith("Message"):
                    continue
                input_value = param.get("Input") or {}
                if input_value.get("InputType") == "USER_INPUT":
                    values = (input_value.get("UserInputValue") or {}).get("Values") or []
                    if values and values[0]:
                        assert "query=" in values[0] and "intent=" in values[0]

        if key in {"02", "03", "04"}:
            wrapped_nodes = [
                node for node in workflow["Nodes"]
                if "_campusflow_runtime_main" in (node.get("CodeExecutorNodeData") or {}).get("Code", "")
            ]
            assert len(wrapped_nodes) >= 2
            primary_output = None
            for node in wrapped_nodes:
                output = run_wrapped_adapter(node["CodeExecutorNodeData"]["Code"], key)
                if node["NodeName"].endswith("Adapter") and "Recovery" not in node["NodeName"]:
                    primary_output = output
            assert primary_output
            if key == "03":
                assert "intent=schedule_day" in primary_output["action0Message"]
                assert "entityName=教师003" in primary_output["action0Message"]

        with zipfile.ZipFile(artifact) as reader:
            assert sorted(reader.namelist()) == sorted([
                f"{workflow['WorkflowID']}_workflow.json", *HEADERS,
            ])
            for workbook_name, expected in HEADERS.items():
                workbook = load_workbook(io.BytesIO(reader.read(workbook_name)), read_only=True)
                header = list(next(workbook.active.iter_rows(values_only=True)))
                assert header == expected, (file_name, workbook_name, header)

    expected_config = json.loads((OUTPUT / "ADP-R4-APP-EXPECTED-CONFIG.json").read_text(encoding="utf-8"))
    assert expected_config["intentRouting"]["schedule_risk_check"] == "03"
    assert expected_config["intentRouting"]["schedule_week"] == "01"
    assert expected_config["pendingWorkflows"][0]["realWidgetId"] is None
    print("test_runtime_binding_r4 passed")


if __name__ == "__main__":
    main()
