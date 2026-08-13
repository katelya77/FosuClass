#!/usr/bin/env python3
import json
import io
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[4]
NATIVE = Path(__file__).resolve().parent
OUTPUT = ROOT / "output" / "competition-adp" / "final"
for key, name in [("02", "02-Classroom-Final.zip"), ("03", "03-Conflict-Final.zip"), ("04", "04-DayPlan-Final.zip")]:
    subprocess.run([sys.executable, str(NATIVE / "validate-campus-widget-artifact.py"),
                    "--key", key, "--artifact", str(OUTPUT / name)], check=True, capture_output=True)
expected = json.loads((OUTPUT / "ADP-App-Expected-Config.json").read_text(encoding="utf-8"))
assert [item["slot"] for item in expected["activeWorkflows"]] == ["01", "02", "03", "04"]
assert expected["routingContract"]["schedule_risk_check"] == "03"
assert len(expected["widgets"]) == 6
assert all("赶场" not in query for query in expected["routerExamples"]["01"])
assert any("赶场" in query for query in expected["routerExamples"]["03"])
with zipfile.ZipFile(OUTPUT / "01-Schedule-Final.zip") as reader:
    workbook = load_workbook(io.BytesIO(reader.read("example_queries.xlsx")), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))[1:]
    examples = [str(row[1] or "") for row in rows]
assert examples and all("冲突" not in query and "赶场" not in query for query in examples)
print("Campus Widget Compiler: PASS (01/02/03/04 + 6 real WidgetIDs)")
