#!/usr/bin/env python3
"""Validate the final Tencent ADP ZIP, not merely compiler source files."""

import argparse
import io
import json
import re
import zipfile
from collections import deque
from pathlib import Path

from openpyxl import load_workbook


NATIVE_DIR = Path(__file__).resolve().parent
TRANSPORT = json.loads((NATIVE_DIR / "adp-transport-contract.json").read_text(encoding="utf-8"))
WIDGET = json.loads((NATIVE_DIR / "schedule-runtime-safe-v3-contract.json").read_text(encoding="utf-8"))
ACTION = json.loads((NATIVE_DIR / "action-contract.json").read_text(encoding="utf-8"))
ROOT_WORKBOOKS = [
    "workflows.xlsx", "parameters.xlsx", "variables.xlsx",
    "example_queries.xlsx", "workflow_references.xlsx",
]
EXPECTED_ROOT = {
    "workflows.xlsx", "parameters.xlsx", "variables.xlsx",
    "example_queries.xlsx", "workflow_references.xlsx",
}


def edge_list(workflow):
    value = workflow.get("Edge", [])
    return json.loads(value) if isinstance(value, str) else value


def all_references(value, owner=None):
    found = []
    if isinstance(value, dict):
        if value.get("InputType") == "REFERENCE_OUTPUT":
            found.append((owner, value.get("Reference", {})))
        for child in value.values():
            found.extend(all_references(child, owner))
    elif isinstance(value, list):
        for child in value:
            found.extend(all_references(child, owner))
    return found


def output_paths(node):
    paths = set()

    def walk(item, prefix=""):
        title = item.get("Title")
        current = f"{prefix}.{title}" if prefix and title else (title or prefix)
        if current:
            paths.add(current)
        child_prefix = current + ("[]" if item.get("Type") in {"ARRAY_OBJECT", "ARRAY_STRING"} else "")
        for child in item.get("Properties", []) or []:
            walk(child, child_prefix)

    for item in node.get("Outputs", []) or []:
        walk(item)
    if node.get("NodeType") == "PARAMETER_EXTRACTOR":
        ui = json.loads(node.get("NodeUI") or "{}")
        content = ui.get("data", {}).get("content", "")
        if isinstance(content, dict):
            for item in content.get("outputs", []):
                if isinstance(item, str):
                    paths.add(item)
        elif isinstance(content, str):
            for item in re.findall(r"[A-Za-z][A-Za-z0-9_]*", content):
                paths.add(f"Output.{item}")
    return paths


def input_reference(node, input_name):
    item = next(
        (value for value in node.get("Inputs", []) or [] if value.get("Name") == input_name),
        None,
    )
    if not item:
        return None
    value = item.get("Input", {})
    if value.get("InputType") != "REFERENCE_OUTPUT":
        return None
    return value.get("Reference", {})


def workbook_rows(data):
    workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    return rows[0] if rows else (), rows[1:] if len(rows) > 1 else []


def run_gate(path):
    checks = []
    failures = []

    def check(name, condition, detail=None):
        if condition:
            checks.append(name)
        else:
            failures.append({"check": name, "detail": detail or "contract mismatch"})

    with zipfile.ZipFile(path, "r") as reader:
        check("ZIP CRC", reader.testzip() is None)
        names = reader.namelist()
        workflow_names = [name for name in names if name.endswith("_workflow.json")]
        check("six-file root contract", len(names) == 6 and all("/" not in name for name in names))
        check("one workflow JSON", len(workflow_names) == 1)
        if len(workflow_names) != 1:
            return checks, failures, None
        workflow_name = workflow_names[0]
        check("required workbook set", set(names) == EXPECTED_ROOT | {workflow_name})
        workflow = json.loads(reader.read(workflow_name))
        workflow_id = workflow.get("WorkflowID")
        workflow_title = workflow.get("WorkflowName")
        check("workflow identity", bool(workflow_id) and bool(workflow_title))
        check("canonical WorkflowID", workflow_id == ACTION["generated"]["workflowId"])
        check("canonical WorkflowName", workflow_title == ACTION["generated"]["workflowName"])
        check("CanvasStructure filename", workflow_name == f"{workflow_id}_workflow.json")

        nodes = workflow.get("Nodes", [])
        node_ids = [node.get("NodeID") for node in nodes]
        node_by_id = {node.get("NodeID"): node for node in nodes}
        check("unique NodeID", len(node_ids) == len(node_by_id) and None not in node_by_id)
        edges = edge_list(workflow)
        edge_pairs = {(edge.get("source"), edge.get("target")) for edge in edges}
        edge_ids = [edge.get("id") for edge in edges]
        edge_handles = [(edge.get("source"), edge.get("target"), edge.get("sourceHandle")) for edge in edges]
        check("unique Edge", len(edge_ids) == len(set(edge_ids)) and len(edge_handles) == len(set(edge_handles)))
        check("Edge endpoints", all(source in node_by_id and target in node_by_id for source, target in edge_pairs))
        declared_pairs = {
            (node["NodeID"], target)
            for node in nodes
            for target in node.get("NextNodeIDs", []) or []
        }
        for node in nodes:
            if node.get("NodeType") == "LOGIC_EVALUATOR":
                for group in node.get("LogicEvaluatorNodeData", {}).get("Group", []):
                    declared_pairs.update((node["NodeID"], target) for target in group.get("NextNodeIDs", []))
        check("NextNodeIDs and branch edges", declared_pairs <= edge_pairs,
              sorted(declared_pairs - edge_pairs))
        invalid_logic_handles = []
        for node in nodes:
            if node.get("NodeType") != "LOGIC_EVALUATOR":
                continue
            groups = node.get("LogicEvaluatorNodeData", {}).get("Group", [])
            ui_content = json.loads(node["NodeUI"]).get("data", {}).get("content", [])
            if not isinstance(ui_content, list) or len(ui_content) != len(groups):
                invalid_logic_handles.append({"node": node["NodeID"], "reason": "UI/group count"})
                continue
            for group, ui_group in zip(groups, ui_content):
                expected_handle = f"{node['NodeID']}.{ui_group.get('id')}-source"
                for target in group.get("NextNodeIDs", []):
                    if not any(edge.get("source") == node["NodeID"]
                               and edge.get("target") == target
                               and edge.get("sourceHandle") == expected_handle for edge in edges):
                        invalid_logic_handles.append({
                            "node": node["NodeID"], "target": target,
                            "expectedSourceHandle": expected_handle,
                        })
        check("Logic branch source handles", not invalid_logic_handles, invalid_logic_handles)

        incoming = {node_id: 0 for node_id in node_ids}
        adjacency = {node_id: [] for node_id in node_ids}
        for source, target in edge_pairs:
            if source in adjacency and target in incoming:
                adjacency[source].append(target)
                incoming[target] += 1
        starts = [node_id for node_id, count in incoming.items() if count == 0]
        visited = set(starts)
        queue = deque(starts)
        while queue:
            current = queue.popleft()
            for target in adjacency[current]:
                if target not in visited:
                    visited.add(target)
                    queue.append(target)
        check("all nodes reachable", len(starts) == 1 and visited == set(node_ids), {
            "starts": starts, "unreachable": sorted(set(node_ids) - visited)})

        refs = []
        for node in nodes:
            refs.extend(all_references(node, node.get("NodeID")))
        check("Reference NodeID", all(ref.get("NodeID") in node_by_id for _, ref in refs))
        invalid_paths = []
        for owner, ref in refs:
            source = node_by_id.get(ref.get("NodeID"))
            path_value = ref.get("JsonPath")
            if not source or not path_value or path_value not in output_paths(source):
                invalid_paths.append({"owner": owner, "source": ref.get("NodeID"), "path": path_value})
        check("Reference JsonPath", not invalid_paths, invalid_paths[:10])

        tool_nodes = {node.get("NodeName"): node for node in nodes if node.get("NodeType") == "TOOL"}
        route_names = {"WEEK": "课表查询-WEEK", "DAY": "课表查询-DAY", "DATE": "课表查询-DATE"}
        schedule_policy = TRANSPORT["workflows"]["01-schedule"]["routes"]
        for route, node_name in route_names.items():
            tool = tool_nodes.get(node_name)
            check(f"{route} Tool exists", tool is not None)
            if not tool:
                continue
            body = tool.get("ToolNodeData", {}).get("Body", [])
            keys = [item.get("ParamName") for item in body]
            policy = schedule_policy[route]
            check(f"{route} request exact keys", keys == policy["requiredKeys"], {
                "expected": policy["requiredKeys"], "actual": keys})
            check(f"{route} forbidden keys omitted", not (set(keys) & set(policy["forbiddenKeys"])))
            check(f"{route} ParamType and SubParams", all(
                item.get("ParamType") == policy["parameters"][item["ParamName"]]["type"]
                and item.get("SubParams") == [] for item in body
            ))
            ui = json.loads(tool["NodeUI"])
            check(f"{route} NodeUI inputs", ui["data"]["content"]["inputs"] == keys)
            api = tool.get("ToolNodeData", {}).get("API", {})
            check(f"{route} Tool identity", (
                api.get("Method") == "POST"
                and api.get("URL", "").rstrip("/").endswith("/api/query_schedule")
                and api.get("CallingMethod") == "NON_STREAMING"
            ), {
                "method": api.get("Method"),
                "urlSuffix": api.get("URL", "").rsplit("/", 2)[-2:],
                "callingMethod": api.get("CallingMethod"),
            })
        week_body = tool_nodes.get("课表查询-WEEK", {}).get("ToolNodeData", {}).get("Body", [])
        week_serialized = json.dumps(week_body, ensure_ascii=False)
        check("WEEK has no weekday token", "weekday" not in week_serialized.lower())
        check("WEEK has no date key", not any(item.get("ParamName") == "date" for item in week_body))

        # A globally valid reference can still be semantically wrong. Each cloned
        # Verify/Adapter branch must consume its own Tool output, never DAY's.
        node_by_name = {node.get("NodeName"): node for node in nodes}
        for route, tool_name in route_names.items():
            tool = node_by_name.get(tool_name)
            for consumer_name in [f"结果核验与呈现-{route}", f"Widget数据适配-Schedule-{route}"]:
                consumer = node_by_name.get(consumer_name)
                reference = input_reference(consumer or {}, "tool_body")
                check(
                    f"{route} {consumer_name} branch-local tool_body",
                    bool(tool and consumer and reference)
                    and reference.get("NodeID") == tool.get("NodeID")
                    and reference.get("JsonPath") == "Output.Body",
                    {
                        "consumer": consumer_name,
                        "expectedToolNodeId": tool.get("NodeID") if tool else None,
                        "actualReference": reference,
                    },
                )

        widget_fields = [item["name"] for item in WIDGET["fields"]]
        widget_types = {item["name"]: item["adpType"] for item in WIDGET["fields"]}
        widget_nodes = [node for node in nodes if node.get("NodeType") == "WIDGET"]
        check("three Schedule Widgets", len(widget_nodes) == 3)
        for widget_node in widget_nodes:
            data = widget_node.get("WidgetNodeData", {})
            params = data.get("WidgetParam", [])
            check(f"WidgetID {widget_node.get('NodeName')}", data.get("WidgetID") == WIDGET["widgetId"])
            check(f"WidgetParam {widget_node.get('NodeName')}", [item.get("ParamName") for item in params] == widget_fields)
            check(f"Widget types {widget_node.get('NodeName')}", all(
                item.get("ParamType") == widget_types[item["ParamName"]]
                and item.get("SubParams") == [] for item in params
            ))
            check(f"Widget ActionType {widget_node.get('NodeName')}", data.get("ActionType") == WIDGET["actionType"])
            ui = json.loads(widget_node["NodeUI"])
            check(f"Widget NodeUI {widget_node.get('NodeName')}", ui["data"]["content"]["inputs"] == widget_fields)

        xlsx = {}
        for name in ROOT_WORKBOOKS:
            headers, rows = workbook_rows(reader.read(name))
            xlsx[name] = {"headers": headers, "rows": rows}
            workflow_index = headers.index("WorkflowId") if "WorkflowId" in headers else -1
            check(f"{name} WorkflowId column", workflow_index >= 0)
            check(f"{name} WorkflowId sync", workflow_index >= 0 and all(
                row[workflow_index] == workflow_id for row in rows
            ))
        workflow_rows = xlsx["workflows.xlsx"]["rows"]
        check("workflows.xlsx one row", len(workflow_rows) == 1)
        if workflow_rows:
            headers = xlsx["workflows.xlsx"]["headers"]
            record = dict(zip(headers, workflow_rows[0]))
            check("workflows.xlsx WorkflowName", record.get("WorkflowName") == workflow_title)
            check("workflows.xlsx CanvasStructure", record.get("CanvasStructure") == workflow_name)

    return checks, failures, {
        "workflowId": workflow_id,
        "workflowName": workflow_title,
        "nodeCount": len(nodes),
        "edgeCount": len(edges),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact", required=True)
    parser.add_argument("--report")
    args = parser.parse_args()
    artifact = Path(args.artifact).resolve()
    checks, failures, summary = run_gate(artifact)
    report = {
        "schema": "fosuclass-adp-validation-report/v1",
        "artifact": artifact.name,
        "status": "PASS" if not failures else "FAIL",
        "checksPassed": len(checks),
        "checksFailed": len(failures),
        "checks": checks,
        "failures": failures,
        "summary": summary,
    }
    if args.report:
        Path(args.report).write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n"
        )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    raise SystemExit(0 if not failures else 1)


if __name__ == "__main__":
    main()
