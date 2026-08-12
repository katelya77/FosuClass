#!/usr/bin/env python3
"""Semantic artifact gate for compiled 02/03/04 native Widget workflows."""

import argparse
import io
import json
import re
import zipfile
from collections import deque
from pathlib import Path

from openpyxl import load_workbook
from campus_widget_compiler import SPECS, CATALOG, REGISTRY, type_name


ROOT = Path(__file__).resolve().parent


def output_paths(node):
    paths = set()
    def walk(item, prefix=""):
        title = item.get("Title")
        current = f"{prefix}.{title}" if prefix and title else (title or prefix)
        if current: paths.add(current)
        for child in item.get("Properties", []): walk(child, current)
    for item in node.get("Outputs", []): walk(item)
    if node.get("NodeType") == "PARAMETER_EXTRACTOR":
        content = json.loads(node.get("NodeUI") or "{}").get("data", {}).get("content", {})
        if isinstance(content, dict):
            paths.update(value for value in content.get("outputs", []) if isinstance(value, str))
        elif isinstance(content, str):
            paths.update(f"Output.{value}" for value in re.findall(r"[A-Za-z][A-Za-z0-9_]*", content))
    return paths


def refs(value):
    result = []
    if isinstance(value, dict):
        if value.get("InputType") == "REFERENCE_OUTPUT": result.append(value.get("Reference", {}))
        for child in value.values(): result.extend(refs(child))
    elif isinstance(value, list):
        for child in value: result.extend(refs(child))
    return result


def gate(path, key):
    spec = SPECS[key]
    checks, failures = [], []
    def check(name, condition, detail=""):
        (checks if condition else failures).append(name if condition else {"check": name, "detail": detail})
    with zipfile.ZipFile(path) as reader:
        names = reader.namelist()
        workflow_names = [name for name in names if name.endswith("_workflow.json")]
        check("six-file import root", len(names) == 6 and len(workflow_names) == 1 and all("/" not in name for name in names))
        workflow = json.loads(reader.read(workflow_names[0]))
        check("real source WorkflowID", workflow.get("WorkflowID") == CATALOG["workflows"][key]["workflowId"])
        check("Final WorkflowName", workflow.get("WorkflowName") == spec["name"])
        nodes = workflow.get("Nodes", [])
        by_id = {node.get("NodeID"): node for node in nodes}
        by_name = {node.get("NodeName"): node for node in nodes}
        check("unique NodeID", len(by_id) == len(nodes) and None not in by_id)
        edges = json.loads(workflow["Edge"]) if isinstance(workflow["Edge"], str) else workflow["Edge"]
        pairs = {(edge.get("source"), edge.get("target")) for edge in edges}
        check("edge endpoints", all(source in by_id and target in by_id for source, target in pairs))
        starts = [node_id for node_id in by_id if not any(target == node_id for _, target in pairs)]
        seen, queue = set(starts), deque(starts)
        while queue:
            current = queue.popleft()
            for source, target in pairs:
                if source == current and target not in seen: seen.add(target); queue.append(target)
        check("all nodes reachable", len(starts) == 1 and seen == set(by_id), {"starts": starts, "unreachable": list(set(by_id) - seen)})
        invalid = []
        for reference in refs(nodes):
            source = by_id.get(reference.get("NodeID"))
            if not source or reference.get("JsonPath") not in output_paths(source): invalid.append(reference)
        check("all Reference outputs valid", not invalid, invalid[:5])

        tool = by_name.get(spec["tool"])
        primary = by_name.get(f"{spec['kind']} Adapter")
        recovery = by_name.get("Choice/Error Adapter")
        router = by_name.get("Native Widget Route")
        check("Tool→primary Adapter chain", bool(tool and primary) and (tool["NodeID"], by_name[spec["verify"]]["NodeID"]) in pairs and
              (by_name[spec["verify"]]["NodeID"], primary["NodeID"]) in pairs)
        for name, node in [("primary", primary), ("recovery", recovery)]:
            item = next((value for value in (node or {}).get("Inputs", []) if value.get("Name") == "tool_body"), {})
            reference = item.get("Input", {}).get("Reference", {})
            check(f"{name} branch-local tool_body", reference == {"NodeID": tool["NodeID"], "JsonPath": "Output.Body"})
        check("recovery router chain", bool(primary and recovery and router) and
              (primary["NodeID"], recovery["NodeID"]) in pairs and (recovery["NodeID"], router["NodeID"]) in pairs)

        widget_expectations = [(spec["kind"], f"{spec['kind']} Adapter"), ("Choice", "Choice/Error Adapter"), ("Error", "Choice/Error Adapter")]
        for kind, adapter_name in widget_expectations:
            record = REGISTRY["widgets"][kind]
            widget = by_name.get(record["name"])
            schema = json.loads((ROOT / "contracts" / kind.lower() / "schema.json").read_text(encoding="utf-8"))
            check(f"{kind} real WidgetID", bool(widget) and widget.get("WidgetNodeData", {}).get("WidgetID") == record["widgetId"])
            params = (widget or {}).get("WidgetNodeData", {}).get("WidgetParam", [])
            check(f"{kind} exported schema fields", [item.get("ParamName") for item in params] == list(schema["properties"]))
            check(f"{kind} exported schema types", all(item.get("ParamType") == type_name(schema["properties"][item["ParamName"]]) for item in params))
            adapter = by_name.get(adapter_name)
            check(f"{kind} Widget references Adapter", all(item.get("Input", {}).get("Reference", {}).get("NodeID") == adapter["NodeID"] for item in params))
        if key == "03":
            missing_widgets = [node for node in nodes if node.get("NodeType") == "WIDGET" and
                               node.get("WidgetNodeData", {}).get("WidgetID") == REGISTRY["widgets"]["Error"]["widgetId"] and
                               "缺少比较对象" in node.get("NodeName", "")]
            check("pre-tool MISSING_PARAM uses real Error Widget", len(missing_widgets) == 1)
            if missing_widgets:
                error_param = next((item for item in missing_widgets[0]["WidgetNodeData"]["WidgetParam"]
                                    if item.get("ParamName") == "error"), {})
                code_param = next((item for item in error_param.get("SubParams", [])
                                   if item.get("ParamName") == "code"), {})
                values = code_param.get("Input", {}).get("UserInputValue", {}).get("Values", [])
                check("pre-tool Error code is MISSING_PARAM", values == ["MISSING_PARAM"])
        headers, rows = list(load_workbook(io.BytesIO(reader.read("example_queries.xlsx")), read_only=True).active.values)[0], \
            list(load_workbook(io.BytesIO(reader.read("example_queries.xlsx")), read_only=True).active.values)[1:]
        examples = [str(row[1]) for row in rows]
        if key == "03": check("risk example belongs to 03", any("赶场" in value for value in examples))
        else: check("risk example excluded from non-03", not any("冲突" in value or "赶场" in value for value in examples), examples)
    return {"schema": "fosuclass-campus-widget-artifact-gate/v1", "artifact": path.name,
            "key": key, "status": "PASS" if not failures else "FAIL", "checksPassed": len(checks),
            "checksFailed": len(failures), "checks": checks, "failures": failures}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--artifact", required=True)
    parser.add_argument("--key", required=True, choices=sorted(SPECS))
    parser.add_argument("--report")
    args = parser.parse_args()
    result = gate(Path(args.artifact), args.key)
    if args.report: Path(args.report).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["status"] == "PASS" else 1)


if __name__ == "__main__": main()
