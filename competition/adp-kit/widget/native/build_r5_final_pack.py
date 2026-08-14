"""Build the judge-ready CampusFlow ADP R5 application closure pack."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from bind_runtime_environment import (
    DEFAULT_R5_OUTPUT,
    EXPECTED_HERO_EXPORT_SHA256,
    EXPECTED_HERO_WIDGET_ID,
    FIXED_ZIP_TIME,
    WORKFLOW_FILES,
    deterministic_xlsx,
    sha256_file,
    workflow_from_zip,
)


ROOT = Path(__file__).resolve().parent
KIT = ROOT.parents[1]
REPO = ROOT.parents[3]
R4_OUTPUT = REPO / "output" / "competition-adp" / "r4"
HERO_ROOT = ROOT / "campus-overview-v1"
DATASET_PATH = KIT / "mock-data" / "competition-demo-v1.json"
GOLDEN_PATH = KIT / "evaluation" / "golden-results.json"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def canonical_json_sha256(value) -> str:
    raw = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()


def workflow_metadata(output_dir: Path) -> list[dict]:
    result = []
    for slot in ("01", "02", "03", "04", "05"):
        path = output_dir / WORKFLOW_FILES[slot]
        if not path.is_file():
            raise AssertionError(f"R5 bound workflow missing: {path.name}")
        workflow, _ = workflow_from_zip(path)
        result.append({
            "slot": slot,
            "workflowId": workflow["WorkflowID"],
            "name": workflow["WorkflowName"],
            "status": "READY_FOR_TENCENT_RUNTIME" if slot == "05" else "TENCENT_WORKFLOW_DEBUG_PASS",
            "artifact": path.name,
            "artifactSha256": sha256_file(path),
        })
    return result


def write_bound_contract(output_dir: Path, hero: dict) -> Path:
    contract = read_json(HERO_ROOT / "contract.json")
    payload = {
        **contract,
        "widgetId": hero["realWidgetId"],
        "realWidgetId": hero["realWidgetId"],
        "bindingStatus": "REAL_EXPORT_BOUND",
        "environment": hero["environment"],
        "sourceType": hero["sourceType"],
        "exportSha256": hero["exportSha256"],
        "sourceTemplateSha256": hero["sourceTemplateSha256"],
        "tencentExportViewSha256": hero["tencentExportViewSha256"],
        "semanticViewSha256": hero["semanticViewSha256"],
        "schemaSha256": hero["schemaSha256"],
        "defaultStateSha256": hero["defaultStateSha256"],
    }
    path = output_dir / "05-Campus-Overview-Contract.json"
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    return path


def write_final_config(output_dir: Path, registry: dict, workflows: list[dict]) -> Path:
    dataset = read_json(DATASET_PATH)
    hero = registry["widgets"]["CampusOverview"]
    config = {
        "schema": "fosuclass-adp-r5-app-final-config/v1",
        "appName": "校园智序 · 小序",
        "productPositioning": "校园教学时空决策智能体",
        "environment": registry["environment"],
        "dataVersion": dataset["meta"]["dataVersion"],
        "dataHash": dataset["dataHash"],
        "runtimeStatus": "PENDING_USER_RUNTIME_E2E",
        "activeWorkflows": workflows,
        "widgets": registry["widgets"],
        "campusOverview": {
            "status": "READY_FOR_TENCENT_RUNTIME",
            "workflow": "05-校园教学态势-R1",
            "artifact": WORKFLOW_FILES["05"],
            "artifactSha256": next(item["artifactSha256"] for item in workflows if item["slot"] == "05"),
            "realWidgetId": hero["realWidgetId"],
            "widgetExportSha256": hero["exportSha256"],
            "schemaSha256": hero["schemaSha256"],
            "semanticViewSha256": hero["semanticViewSha256"],
            "dataVersion": dataset["meta"]["dataVersion"],
            "dataHash": dataset["dataHash"],
        },
        "intentRouting": {
            "schedule_day": "01",
            "schedule_week": "01",
            "schedule_choose_day": "01",
            "classroom_find": "02",
            "schedule_risk_check": "03",
            "day_plan": "04",
            "campus_overview": "05",
        },
        "routerExamples": {
            "01": ["教师003第1周周一有什么课"],
            "02": ["校区A第1周周一第5-6节有哪些空教室"],
            "03": ["教师003第1周周一跨校区赶不赶得上"],
            "04": ["帮我安排2026-09-04的一天"],
            "05": [
                "未来四周校园教学情况怎么样",
                "从8月25日开始看看未来几周校园教学运行情况",
                "最近哪一天最忙",
                "哪个校区下午教室最紧张",
            ],
        },
        "routingBoundaries": {
            "overallCampusPressure": "05",
            "specificRoomAvailability": "02",
            "singleEntitySchedule": "01",
            "singleTeacherRisk": "03",
            "personalDayPlan": "04",
        },
        "excludedFromRouting": [
            "01-多维课表查询",
            "02-空教室规划",
            "03-课程冲突比较",
            "04-今日校园计划",
            "01-多维课表查询-Final_9332",
            "00-节点格式种子-勿启用",
            "所有旧 RuntimeSafe / Final 中间版本",
        ],
    }
    path = output_dir / "ADP-R5-APP-FINAL-CONFIG.json"
    path.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    return path


def r5_cases() -> list[list]:
    cases = [
        ("r5-app-001", "router-overview-vs-schedule", "未来四周校园教学情况怎么样", "05-校园教学态势-R1", "get_campus_teaching_overview", "CampusOverview", "campus_overview", "overall overview must not route 01"),
        ("r5-app-002", "router-overview-vs-classroom", "哪个校区整体资源最紧张", "05-校园教学态势-R1", "get_campus_teaching_overview", "CampusOverview", "campus_overview", "overall pressure must not route 02"),
        ("r5-app-003", "05-widget-real-binding", "运行校园教学态势", "05-校园教学态势-R1", "binding gate", "CampusOverview", "runtime_binding", EXPECTED_HERO_WIDGET_ID),
        ("r5-app-004", "05-preparation-period", "8月25日至30日有几节课", "05-校园教学态势-R1", "get_campus_teaching_overview", "CampusOverview", "campus_overview", "preparation lessonCount=0"),
        ("r5-app-005", "05-matrix", "查看四周工作日负载", "05-校园教学态势-R1", "get_campus_teaching_overview", "CampusOverview", "campus_overview", "W1 7,6,7,5,6; W4 7,6,7,5,7"),
        ("r5-app-006", "05-teacher-load", "未来四周教师负载最高是谁", "05-校园教学态势-R1", "get_campus_teaching_overview", "CampusOverview", "schedule_week", "教师002; 24 occurrences"),
        ("r5-app-007", "05-risk-summary", "校园整体教学风险如何", "05-校园教学态势-R1", "get_campus_teaching_overview", "CampusOverview", "schedule_risk_check", "conflict=0;rush=8;continuous=24"),
        ("r5-app-008", "05-action-to-01", "点击查看第1周", "01-多维课表查询-R3", "query_schedule", "Schedule", "schedule_week", "week=1"),
        ("r5-app-009", "05-action-to-02", "点击查空教室并补充周一第5-6节", "02-空教室规划-R3", "find_available_classrooms", "Classroom", "classroom_find", "week=1;weekday=1;period=5-6"),
        ("r5-app-010", "05-action-to-03", "点击检查风险", "03-课程冲突比较-R3", "compare_schedules", "Conflict", "schedule_risk_check", "教师003 self-compare"),
        ("r5-app-011", "03-action-to-01", "点击查看当天课表", "01-多维课表查询-R3", "query_schedule", "Schedule", "schedule_day", "教师003;week=1;weekday=1"),
        ("r5-app-012", "01-to-04-context", "帮我把当天空档安排成连续自习", "04-今日校园计划-R3", "generate_day_plan", "DayPlan", "day_plan", "confirmed date preserved"),
        ("r5-app-013", "old-workflow-not-routed", "教师003第1周周一的课", "01-多维课表查询-R3", "query_schedule", "Schedule", "schedule_day", "historical baseline excluded"),
        ("r5-app-014", "semantic-widget-normalization", "Tencent formatter-only Widget diff", "05-校园教学态势-R1", "semantic binding gate", "CampusOverview", "runtime_binding", "raw hash differs; semantic hash equal"),
        ("r5-app-015", "real-widget-id-gate", "05 import resource check", "05-校园教学态势-R1", "binding gate", "CampusOverview", "runtime_binding", "32 hex, unique, real export"),
    ]
    return [[
        case_id, category, user_input, "[]", reference, workflow, tool, widget,
        True, action, "rule/code", "anonymousOnly + noFabrication + environmentBinding",
    ] for case_id, category, user_input, workflow, tool, widget, action, reference in cases]


def build_evaluation(output_dir: Path) -> tuple[Path, int]:
    source = R4_OUTPUT / "ADP-R4-EVALUATION.xlsx"
    workbook = load_workbook(source, read_only=True, data_only=False)
    rows = list(workbook.active.iter_rows(values_only=True))
    if len(rows) != 93:
        raise AssertionError("R4 evaluation must contain header + 92 preserved cases")
    headers = list(rows[0])
    data_rows = [list(row) for row in rows[1:]]
    data_rows.extend(r5_cases())
    path = output_dir / "ADP-R5-FINAL-EVALUATION.xlsx"
    path.write_bytes(deterministic_xlsx(headers, data_rows))
    return path, len(data_rows)


def write_docs(output_dir: Path, registry: dict) -> list[Path]:
    widgets = registry["widgets"]
    activation = """# ADP R5 Application Activation Matrix

Workflow debug PASS 只证明单个画布可启动，不证明 Agent Router 使用了该版本。最终应用必须按下表排除新旧路由竞争。

| 状态 | Workflow | 用途 |
|---|---|---|
| ENABLE | 01-多维课表查询-R3 | 单实体课表与占用 |
| ENABLE | 02-空教室规划-R3 | 具体时间条件下的空间资源决策 |
| ENABLE | 03-课程冲突比较-R3 | 两实体冲突与教师 self-compare 赶场 |
| ENABLE | 04-今日校园计划-R3 | 已确认日期的个人行动规划 |
| ENABLE | 05-校园教学态势-R1 | 校园整体态势与 01-04 调度台 |
| DISABLE / EXCLUDE | 01/02/03/04 历史 baseline | 避免同意图多版本竞争 |
| DISABLE / EXCLUDE | 01-多维课表查询-Final_9332 | 历史结构证据，仅回滚使用 |
| DISABLE / EXCLUDE | 00-节点格式种子-勿启用 | 研发 seed，不参与应用路由 |
| DISABLE / EXCLUDE | 所有旧 RuntimeSafe / Final 中间版本 | 避免 old/new routing collision |

不要删除历史 Workflow；只关闭应用引用和 Router 示例，保留证据与回滚能力。
"""
    checklist = f"""# ADP R5 Runtime E2E Checklist

状态必须分开记录：`LOCAL_CONTRACT_PASS`、`TENCENT_WORKFLOW_DEBUG_PASS`、`TENCENT_APPLICATION_E2E_PASS`。

| 场景 | 输入 | Workflow | Tool | Widget / ID | 当前状态 | 必须出现 |
|---|---|---|---|---|---|---|
| Schedule | 教师003第1周周一的课 | 01-R3 | query_schedule | Schedule `{widgets['Schedule']['realWidgetId']}` | TENCENT_WORKFLOW_DEBUG_PASS | 2门课、verified、competition-demo-v1 |
| Classroom | 校区A第1周周一第5-6节有哪些空教室 | 02-R3 | find_available_classrooms | Classroom `{widgets['Classroom']['realWidgetId']}` | TENCENT_WORKFLOW_DEBUG_PASS | 筛选条件、原生卡 |
| Conflict | 教师003第1周周一跨校区赶不赶得上 | 03-R3 | compare_schedules self-compare | Conflict `{widgets['Conflict']['realWidgetId']}` | TENCENT_WORKFLOW_DEBUG_PASS | 冲突0、赶场1、不要求第二对象 |
| DayPlan | 帮我安排2026-09-04的一天 | 04-R3 | generate_day_plan | DayPlan `{widgets['DayPlan']['realWidgetId']}` | TENCENT_WORKFLOW_DEBUG_PASS | 课程、空档、自习建议 |
| Choice | 稳定歧义 fixture | 原任务 Recovery | resolve_entity | Choice `{widgets['Choice']['realWidgetId']}` | PENDING_USER_APPLICATION_E2E | 候选确认后回原任务 |
| Error | 查空教室（缺日期/节次） | 02 Recovery | MISSING_PARAM | Error `{widgets['Error']['realWidgetId']}` | PENDING_USER_APPLICATION_E2E | 中文恢复；不要求 queryId |
| Hero | 从8月25日开始看看未来几周校园教学运行情况 | 05-R1 | get_campus_teaching_overview | CampusOverview `{widgets['CampusOverview']['realWidgetId']}` | LOCAL_CONTRACT_PASS | 准备期0课、126次、4×5、教师/资源/风险、已核验 |

## Hero application chain

1. Hero「查空教室」必须发送 `intent=classroom_find|week=1` 并进入 02；补充周一第5-6节。
2. 「检查风险」必须发送 `intent=schedule_risk_check|entityType=teacher|entityName=教师003|week=1` 并进入 03 self-compare。
3. Conflict「查看当天课表」必须发送 `intent=schedule_day` 并进入 01，保留教师003、第1周、周一。
4. `帮我把当天空档安排成连续自习` 必须进入 04 并沿用已确认日期。
5. 腾讯应用级全链完成前，整体状态保持 `PENDING_USER_E2E`。
"""
    judge = """# ADP R5 Judge Demo（2～4 分钟）

核心叙事：校园教学信息分散在课表、教师、教室与个人安排中。小序先用确定性工具形成校园教学态势，再从资源压力钻取到教室、教师风险和个人计划；全链共享同一匿名数据版本与可核验证据。

| 演示输入 / 点击 | 预期 UI | 对评委说的一句话 | 下一步 |
|---|---|---|---|
| `从8月25日开始看看未来几周校园教学运行情况` | 05 Hero：准备期0课、4周、126次、4×5矩阵、资源/教师/风险、已核验 | 小序不是自己“数课”，全部指标由确定性 CampusTools 从同一事实模型派生。 | 点击「查空教室」 |
| 补充 `第1周周一第5-6节，校区A` | 02 空教室原生卡 | 从校园整体压力进入具体空间决策，事实版本没有切换。 | 点击/输入「检查教师003风险」 |
| `检查教师003第1周周一跨校区赶场` | 03：冲突0、赶场1 | 单教师采用 self-compare，不伪造第二对象。 | 点击「查看当天课表」 |
| Conflict「查看当天课表」 | 01：教师003、W1、周一、2门课 | Action Protocol 携带 intent 与上下文，不靠按钮文字猜路由。 | 输入连续自习请求 |
| `帮我把当天空档安排成连续自习` | 04 DayPlan | 同一日期上下文继续转化为可执行的个人安排。 | 结束 |

收束：小序是一套将分散校园教学事实转化为可核验、可解释、可执行决策的校园教学时空决策智能体，而不是普通聊天机器人或课表展示 Demo。
"""
    rollback = """# R5 Rollback

1. 若 05 Runtime 失败，仅从应用 Router 停用 `05-校园教学态势-R1`；不要删除真实 Widget 或历史 Workflow。
2. 01～04 R3 ZIP 与 R4 实机结构保持 byte-stable，可继续作为稳定基线。
3. 若应用路由竞争，按 Activation Matrix 关闭旧 baseline / Final / RuntimeSafe 中间版本引用。
4. 回滚不涉及 production 数据、VPS、CloudBase、微信小程序、main 或正式 ADP 发布。
"""
    payloads = {
        "ADP-R5-APP-ACTIVATION-MATRIX.md": activation,
        "ADP-R5-RUNTIME-E2E-CHECKLIST.md": checklist,
        "ADP-R5-JUDGE-DEMO.md": judge,
        "ROLLBACK.md": rollback,
    }
    paths = []
    for name, content in payloads.items():
        path = output_dir / name
        path.write_text(content, encoding="utf-8", newline="\n")
        paths.append(path)
    return paths


def write_bundle(output_dir: Path, names: list[str]) -> Path:
    target = output_dir / "CampusFlow-ADP-R5-Final-Pack.zip"
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as writer:
        for name in sorted(names):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            writer.writestr(info, (output_dir / name).read_bytes())
    return target


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", default=str(DEFAULT_R5_OUTPUT))
    args = parser.parse_args()
    output_dir = Path(args.output_dir).resolve()
    registry_path = output_dir / "widget-registry.runtime.json"
    registry = read_json(registry_path)
    if set(registry.get("widgets", {})) != {"Schedule", "Classroom", "Conflict", "DayPlan", "Choice", "Error", "CampusOverview"}:
        raise AssertionError("R5 registry must contain exactly seven logical Widget kinds")
    hero = registry["widgets"]["CampusOverview"]
    if hero["realWidgetId"] != EXPECTED_HERO_WIDGET_ID or hero["exportSha256"] != EXPECTED_HERO_EXPORT_SHA256:
        raise AssertionError("R5 CampusOverview is not bound to the supplied Tencent export")
    if sha256_file(output_dir / "05-Tencent-Widget-Evidence.widget") != EXPECTED_HERO_EXPORT_SHA256:
        raise AssertionError("R5 Tencent Widget evidence is missing or changed")
    workflows = workflow_metadata(output_dir)
    contract_path = write_bound_contract(output_dir, hero)
    config_path = write_final_config(output_dir, registry, workflows)
    evaluation_path, evaluation_count = build_evaluation(output_dir)
    doc_paths = write_docs(output_dir, registry)
    golden = read_json(GOLDEN_PATH)
    golden_count = len(golden) if isinstance(golden, list) else golden.get("count")
    artifact_names = [
        *[WORKFLOW_FILES[slot] for slot in ("01", "02", "03", "04", "05")],
        "05-Tencent-Widget-Evidence.widget",
        "widget-registry.runtime.json",
        contract_path.name,
        config_path.name,
        evaluation_path.name,
        *[path.name for path in doc_paths],
    ]
    manifest = {
        "schema": "fosuclass-adp-r5-final-pack/v1",
        "environment": registry["environment"],
        "dataVersion": read_json(DATASET_PATH)["meta"]["dataVersion"],
        "dataHash": read_json(DATASET_PATH)["dataHash"],
        "runtimeBinding": "REAL_EXPORT_BOUND",
        "tencentRuntimeStatus": "PENDING_USER_E2E",
        "workflows": workflows,
        "evaluation": {"baselineCases": 80, "r4HeroCases": 12, "r5ApplicationCases": 15, "totalCases": evaluation_count},
        "golden": {"expected": 33, "actual": golden_count},
        "artifacts": [{"file": name, "sha256": sha256_file(output_dir / name)} for name in artifact_names],
    }
    manifest_path = output_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    artifact_names.append(manifest_path.name)
    sums_path = output_dir / "SHA256SUMS.txt"
    sums_path.write_text(
        "".join(f"{sha256_file(output_dir / name)}  {name}\n" for name in sorted(artifact_names)),
        encoding="utf-8", newline="\n",
    )
    artifact_names.append(sums_path.name)
    bundle = write_bundle(output_dir, artifact_names)
    print(f"ADP R5 final pack: PASS ({evaluation_count} evaluation cases)\n{bundle}\nSHA256={sha256_file(bundle)}")


if __name__ == "__main__":
    main()
