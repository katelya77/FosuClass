import hashlib
import io
import json
import runpy
import subprocess
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[3]
OUTPUT = REPO / "output" / "competition-adp" / "r5"
R4_OUTPUT = REPO / "output" / "competition-adp" / "r4"

from bind_runtime_environment import (  # noqa: E402
    EXPECTED_HERO_WIDGET_ID,
    WORKFLOW_FILES,
    workflow_from_zip,
)


EXPECTED_IDS = {
    "Schedule": "23fbc659efe3482fab588d754e4420a4",
    "Classroom": "d781773ab49c4b15a5a4b99f0e28a748",
    "Conflict": "86e9ea7eb9e0458cb69a5d08aa767163",
    "DayPlan": "9e334031b8bc47e9ab349244f9bca421",
    "Choice": "b2726b1796294285b9f2daffc6d75dc7",
    "Error": "906df4c8d24548fabb8bd8a080de76d2",
    "CampusOverview": EXPECTED_HERO_WIDGET_ID,
}
HEADERS = {
    "workflows.xlsx": ["WorkflowId", "WorkflowName", "WorkflowDescription", "CanvasStructure"],
    "parameters.xlsx": ["WorkflowId", "WorkflowNodeId", "WorkflowNodeName", "ParameterId", "ParameterName", "ParameterDescription", "ParameterType", "ParameterCorrectExample", "ParameterIncorrectExample", "ParameterParentId", "ParameterRequiredStatus"],
    "example_queries.xlsx": ["WorkflowId", "ExampleQueryContent"],
    "variables.xlsx": ["WorkflowId", "VariableId", "VariableName", "VariableDescription", "VariableType", "VariableDefaultValue", "VariableDefaultValueFileName", "ParameterType"],
    "workflow_references.xlsx": ["WorkflowId", "WorkflowNodeId", "ReferenceWorkflowId", "ReferenceWorkflowName"],
}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def call_overview_tool():
    script = (
        "const {callTool}=require('./competition/adp-kit/mcp/campus-tools-mcp/src/tools');"
        "console.log(JSON.stringify(callTool('get_campus_teaching_overview',{})));"
    )
    result = subprocess.run(
        ["node", "-e", script], cwd=REPO, capture_output=True, text=True,
        encoding="utf-8", check=True,
    )
    return json.loads(result.stdout)


def protocol(message):
    assert "【小序ActionV2】" in message
    natural, encoded = message.split("【小序ActionV2】", 1)
    fields = {}
    for part in encoded.strip().split("|"):
        key, value = part.split("=", 1)
        fields[key] = value
    assert natural.strip() == fields["query"]
    return fields


def route(config, query, intent=None):
    if intent:
        return config["intentRouting"][intent]
    if any(value in query for value in ("未来四周", "未来几周", "最近哪一天最忙", "整体资源", "整体压力", "校园整体")):
        return "05"
    if any(value in query for value in ("冲突", "赶不赶得上", "赶场")):
        return "03"
    if "空教室" in query or ("教室" in query and "有哪些" in query):
        return "02"
    if any(value in query for value in ("帮我安排", "一天", "连续自习")):
        return "04"
    return "01"


def main():
    registry = json.loads((OUTPUT / "widget-registry.runtime.json").read_text(encoding="utf-8"))
    assert registry["environment"] == "competition-current"
    assert {key: value["realWidgetId"] for key, value in registry["widgets"].items()} == EXPECTED_IDS
    ids = [record["realWidgetId"] for record in registry["widgets"].values()]
    assert len(ids) == len(set(ids)) == 7
    assert all(len(value) == 32 and all(char in "0123456789abcdef" for char in value) for value in ids)

    for slot in ("01", "02", "03", "04"):
        assert sha256(OUTPUT / WORKFLOW_FILES[slot]) == sha256(R4_OUTPUT / WORKFLOW_FILES[slot])

    artifact = OUTPUT / WORKFLOW_FILES["05"]
    workflow, workflow_json_name = workflow_from_zip(artifact)
    assert workflow["WorkflowName"] == "05-校园教学态势-R1"
    assert not any(node.get("NodeType") in {"LLM", "PARAMETER_EXTRACTOR"} for node in workflow["Nodes"])
    names = [node["NodeName"] for node in workflow["Nodes"]]
    assert names == [
        "开始", "校园教学态势分析", "CampusOverview Result Guard",
        "CampusOverview Adapter", "CampusOverview Recovery Adapter",
        "CampusOverview Success/Error Route", "小序-校园教学态势-RuntimeSafe-V1",
        "小序-任务恢复-RuntimeSafe-V3", "结束",
    ]
    tool = next(node for node in workflow["Nodes"] if node["NodeName"] == "校园教学态势分析")
    assert tool["ToolNodeData"]["API"]["URL"].endswith("/api/get_campus_teaching_overview")
    assert {param["ParamName"]: param["Input"]["UserInputValue"]["Values"][0] for param in tool["ToolNodeData"]["Body"]} == {
        "windowStart": "2026-08-25", "teachingStart": "2026-08-31", "windowEnd": "2026-09-27",
    }
    guard = next(node for node in workflow["Nodes"] if node["NodeName"] == "CampusOverview Result Guard")
    assert "sha1:fefef4bf425b" in guard["CodeExecutorNodeData"]["Code"]
    assert "lessonCount\") == 0" in guard["CodeExecutorNodeData"]["Code"]
    widgets = [node for node in workflow["Nodes"] if node["NodeType"] == "WIDGET"]
    assert {(node["WidgetNodeData"]["WidgetID"], node["NodeName"]) for node in widgets} == {
        (EXPECTED_IDS["CampusOverview"], "小序-校园教学态势-RuntimeSafe-V1"),
        (EXPECTED_IDS["Error"], "小序-任务恢复-RuntimeSafe-V3"),
    }
    assert all(
        param["ParamType"] in {"STRING", "INT", "FLOAT", "BOOL"}
        for node in widgets for param in node["WidgetNodeData"]["WidgetParam"]
    )
    hero_widget = next(node for node in widgets if node["WidgetNodeData"]["WidgetID"] == EXPECTED_IDS["CampusOverview"])
    contract = json.loads((ROOT / "campus-overview-v1" / "contract.json").read_text(encoding="utf-8"))
    assert [param["ParamName"] for param in hero_widget["WidgetNodeData"]["WidgetParam"]] == contract["fields"]

    with zipfile.ZipFile(artifact) as reader:
        assert set(reader.namelist()) == {workflow_json_name, *HEADERS}
        for name, expected in HEADERS.items():
            workbook = load_workbook(io.BytesIO(reader.read(name)), read_only=True)
            assert list(next(workbook.active.iter_rows(values_only=True))) == expected

    overview = call_overview_tool()
    namespace = {}
    exec(guard["CodeExecutorNodeData"]["Code"], namespace)
    guarded = namespace["main"]({"tool_body": overview})["Body"]
    assert guarded["success"] is True
    hero_adapter = runpy.run_path(str(ROOT / "campus-overview-v1" / "adapter.py"))["main"]
    hero_output = hero_adapter({"tool_body": guarded})
    assert hero_output["lessonCountText"] == "126 次课程"
    assert [hero_output[f"w{week}{day}Text"] for week in range(1, 5) for day in ("Mon", "Tue", "Wed", "Thu", "Fri")] == [
        "7", "6", "7", "5", "6", "7", "6", "7", "5", "7",
        "7", "6", "7", "5", "6", "7", "6", "7", "5", "7",
    ]
    action0 = protocol(hero_output["action0Message"])
    action1 = protocol(hero_output["action1Message"])
    action2 = protocol(hero_output["action2Message"])
    action3 = protocol(hero_output["action3Message"])
    assert action0 == {"query": "查看第1周校园课表", "intent": "schedule_week", "week": "1"}
    assert action1["intent"] == "classroom_find" and action1["week"] == "1"
    assert action2["intent"] == "schedule_week" and action2["entityType"] == "teacher" and action2["entityName"] == "教师002"
    assert action3["intent"] == "schedule_risk_check" and action3["entityName"] == "教师003" and action3["week"] == "1"

    config = json.loads((OUTPUT / "ADP-R5-APP-FINAL-CONFIG.json").read_text(encoding="utf-8"))
    cases = {
        "教师003第1周周一有什么课": "01",
        "校区A第1周周一第5-6节有哪些空教室": "02",
        "教师003第1周周一跨校区赶不赶得上": "03",
        "帮我安排2026-09-04的一天": "04",
        "未来四周校园教学情况怎么样": "05",
        "从8月25日开始看看未来几周校园教学运行情况": "05",
        "最近哪一天最忙": "05",
        "哪个校区整体资源最紧张": "05",
    }
    for query, expected in cases.items():
        assert route(config, query) == expected, query
    for intent, expected in config["intentRouting"].items():
        assert route(config, "ignored context", intent) == expected
    assert route(config, action1["query"], action1["intent"]) == "02"
    assert route(config, action3["query"], action3["intent"]) == "03"
    assert "01-多维课表查询" in config["excludedFromRouting"]
    assert "00-节点格式种子-勿启用" in config["excludedFromRouting"]

    conflict_workflow, _ = workflow_from_zip(OUTPUT / WORKFLOW_FILES["03"])
    conflict_adapter = next(node for node in conflict_workflow["Nodes"] if node["NodeName"] == "Conflict Adapter")
    conflict_body = {
        "success": True, "dataVersion": "competition-demo-v1", "queryId": "q-r5",
        "evidence": {"verified": True}, "query": {"week": 1, "weekday": 1},
        "items": [], "rushWarnings": [],
        "summary": {"selfCompare": True, "conflictCount": 0, "rushWarningCount": 1},
        "compared": [{"type": "teacher", "name": "教师003"}, {"type": "teacher", "name": "教师003"}],
    }
    conflict_ns = {}
    exec(conflict_adapter["CodeExecutorNodeData"]["Code"], conflict_ns)
    conflict_output = conflict_ns["main"]({"tool_body": conflict_body})
    schedule_handoff = protocol(conflict_output["action0Message"])
    assert schedule_handoff["intent"] == "schedule_day"
    assert schedule_handoff["entityName"] == "教师003" and schedule_handoff["week"] == "1" and schedule_handoff["weekday"] == "1"
    assert route(config, schedule_handoff["query"], schedule_handoff["intent"]) == "01"
    assert route(config, "帮我把当天空档安排成连续自习") == "04"

    activation = (OUTPUT / "ADP-R5-APP-ACTIVATION-MATRIX.md").read_text(encoding="utf-8")
    for name in ("01-多维课表查询-R3", "02-空教室规划-R3", "03-课程冲突比较-R3", "04-今日校园计划-R3", "05-校园教学态势-R1"):
        assert f"ENABLE | {name}" in activation
    assert "00-节点格式种子-勿启用" in activation

    workbook = load_workbook(OUTPUT / "ADP-R5-FINAL-EVALUATION.xlsx", read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert len(rows) == 108
    case_ids = {row[0] for row in rows[1:]}
    assert sum(str(value).startswith("eval-") for value in case_ids) == 80
    assert sum(str(value).startswith("r4-hero-") for value in case_ids) == 12
    assert sum(str(value).startswith("r5-app-") for value in case_ids) == 15
    assert len(case_ids) == 107
    print("test_application_convergence_r5 passed")


if __name__ == "__main__":
    main()
