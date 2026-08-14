import base64
import importlib.util
import json
import subprocess
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[3]
OUTPUT = REPO / "output" / "competition-adp" / "r4"
HERO = ROOT / "campus-overview-v1"


def load_adapter():
    spec = importlib.util.spec_from_file_location("campus_overview_adapter", HERO / "adapter.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def tool_call(params=None):
    script = (
        "const {callTool}=require('./competition/adp-kit/mcp/campus-tools-mcp/src/tools');"
        f"console.log(JSON.stringify(callTool('get_campus_teaching_overview',{json.dumps(params or {})})));"
    )
    result = subprocess.run(["node", "-e", script], cwd=REPO, capture_output=True, text=True, encoding="utf-8", check=True)
    return json.loads(result.stdout)


def main():
    result = tool_call()
    assert result["success"] is True
    assert result["dataVersion"] == "competition-demo-v1"
    assert result["evidence"]["verified"] is True
    assert result["evidence"]["dataHash"] == "sha1:fefef4bf425b"
    overview = result["items"][0]
    assert overview["window"]["preparationPeriod"] == {
        "startDate": "2026-08-25", "endDate": "2026-08-30", "lessonCount": 0,
    }
    assert overview["summary"]["weekCount"] == 4
    assert overview["summary"]["lessonOccurrences"] == 126
    assert [[day["lessonCount"] for day in week["days"]] for week in overview["matrix"]] == [
        [7, 6, 7, 5, 6], [7, 6, 7, 5, 7], [7, 6, 7, 5, 6], [7, 6, 7, 5, 7],
    ]
    assert overview["teacherLoadTop"][0]["teacherName"] == "教师002"
    assert overview["teacherLoadTop"][0]["lessonOccurrences"] == 24
    assert overview["risks"]["conflictCount"] == 0
    assert overview["risks"]["rushCount"] == 8
    assert overview["risks"]["continuousLoadCount"] == 24
    assert result["actions"][2]["intent"] == "schedule_risk_check"

    invalid = tool_call({"windowStart": "2026-08-24"})
    assert invalid["success"] is False
    assert invalid["error"]["code"] == "INVALID_PARAM"

    adapter = load_adapter()
    output = adapter.main({"tool_body": result})
    assert output["w1MonText"] == "7"
    assert output["w4FriText"] == "7"
    assert output["lessonCountText"] == "126 次课程"
    assert output["teacher0Name"] == "教师002"
    assert "intent=schedule_week" in output["action0Message"]
    assert "intent=classroom_find" in output["action1Message"]
    assert "entityName=教师002" in output["action2Message"]
    assert "intent=schedule_risk_check" in output["action3Message"]
    assert "entityName=教师003" in output["action3Message"]

    widget = json.loads(next(OUTPUT.glob("05-*-R1-Pilot.widget")).read_text(encoding="utf-8"))
    inner = json.loads(base64.b64decode(widget["encodedWidget"]).decode("utf-8"))
    contract = json.loads((HERO / "contract.json").read_text(encoding="utf-8"))
    assert inner["id"] == ""
    assert inner["name"] == "小序-校园教学态势-RuntimeSafe-V1"
    assert list(inner["defaultState"]) == contract["fields"]
    assert list(widget["jsonSchema"]["properties"]) == contract["fields"]
    assert set(widget["jsonSchema"]["required"]) == set(contract["fields"])
    assert all(value["type"] == "string" for value in widget["jsonSchema"]["properties"].values())
    assert "Array<Object>" not in inner["schema"]
    assert "Chart" not in inner["view"]

    workbook = load_workbook(OUTPUT / "ADP-R4-EVALUATION.xlsx", read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert len(rows) == 93
    assert len(rows[0]) == 12
    assert sum(1 for row in rows[1:] if str(row[0]).startswith("eval-")) == 80
    assert sum(1 for row in rows[1:] if str(row[0]).startswith("r4-hero-")) == 12
    print("test_campus_overview_r4 passed")


if __name__ == "__main__":
    main()
