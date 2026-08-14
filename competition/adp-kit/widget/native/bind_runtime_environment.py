"""Bind Tencent ADP workflow exports to Widget resources from one environment.

The logical registry remains stable.  This compiler creates an environment-scoped
runtime registry and refuses empty, stale, synthetic, or schema-drifted bindings.
"""

from __future__ import annotations

import argparse
import base64
import datetime
import hashlib
import io
import json
import re
import tempfile
import uuid
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[3]
DEFAULT_OUTPUT = REPO / "output" / "competition-adp" / "r4"
LOGICAL_REGISTRY_PATH = ROOT / "widget-registry.json"
FIXED_ZIP_TIME = (2026, 8, 14, 0, 0, 0)
R4_NAMESPACE = uuid.UUID("d56d7b28-60b9-48a2-9d07-bf2ff7bd1414")
ROOT_FILES = [
    "workflows.xlsx",
    "parameters.xlsx",
    "example_queries.xlsx",
    "variables.xlsx",
    "workflow_references.xlsx",
]
RUNTIME_KINDS = ("Schedule", "Classroom", "Conflict", "DayPlan", "Choice", "Error")
WIDGET_NAME_MARKERS = {
    "Classroom": "空教室票据",
    "Conflict": "冲突赶场票据",
    "DayPlan": "今日校园计划",
    "Choice": "候选确认",
    "Error": "任务恢复",
}
WORKFLOW_NAMES = {
    "01": "01-多维课表查询-R3",
    "02": "02-空教室规划-R3",
    "03": "03-课程冲突比较-R3",
    "04": "04-今日校园计划-R3",
}
WORKFLOW_FILES = {
    "01": "01-多维课表查询-R3-Bound.zip",
    "02": "02-空教室规划-R3-Bound.zip",
    "03": "03-课程冲突比较-R3-Bound.zip",
    "04": "04-今日校园计划-R3-Bound.zip",
}
WORKFLOW_SOURCE_PATTERNS = {
    "01": re.compile(r"^01-多维课表查询-Final(?:_|$)"),
    "02": re.compile(r"^02-空教室规划-RuntimeSafeV3$"),
    "03": re.compile(r"^03-课程冲突比较-RuntimeSafeV3$"),
    "04": re.compile(r"^04-今日校园计划-RuntimeSafeV3$"),
}
WORKFLOW_EXAMPLES = {
    "01": ["教师003第1周周一的课", "A1-101第1周周三的占用"],
    "02": ["校区A第1周周一第5-6节有哪些空教室"],
    "03": ["教师003第1周周一跨校区赶不赶得上", "比较2025级A班和2025级B班第1周周五下午的课程冲突"],
    "04": ["帮我安排2026-09-04的一天"],
}


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def canonical_json(value) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def exported_at(path: Path) -> str:
    stamp = datetime.datetime.fromtimestamp(path.stat().st_mtime, datetime.timezone.utc)
    return stamp.replace(microsecond=0).isoformat().replace("+00:00", "Z")


def parse_widget(path: Path) -> dict:
    raw = path.read_bytes()
    outer = json.loads(raw.decode("utf-8-sig"))
    encoded = outer.get("encodedWidget")
    if not isinstance(encoded, str) or not encoded:
        raise AssertionError(f"{path.name}: encodedWidget missing")
    try:
        inner = json.loads(base64.b64decode(encoded).decode("utf-8"))
    except Exception as error:
        raise AssertionError(f"{path.name}: encodedWidget invalid") from error
    name = str(inner.get("name") or outer.get("name") or "").strip()
    kind = next((key for key, marker in WIDGET_NAME_MARKERS.items() if marker in name), None)
    if not kind:
        raise AssertionError(f"{path.name}: unknown logical Widget kind")
    real_id = str(inner.get("id") or "").strip()
    if not re.fullmatch(r"[0-9a-f]{32}", real_id):
        raise AssertionError(f"{path.name}: real WidgetID is empty or invalid")
    default_state = inner.get("defaultState")
    json_schema = outer.get("jsonSchema")
    if not isinstance(default_state, dict) or not isinstance(json_schema, dict):
        raise AssertionError(f"{path.name}: defaultState/jsonSchema missing")
    schema_fields = list((json_schema.get("properties") or {}).keys())
    default_fields = list(default_state.keys())
    if set(schema_fields) != set(default_fields):
        raise AssertionError(f"{path.name}: DefaultState and JSON Schema field drift")
    if json_schema.get("type") != "object" or json_schema.get("additionalProperties") is not False:
        raise AssertionError(f"{path.name}: runtime Widget schema must be strict object")
    if set(json_schema.get("required") or []) != set(schema_fields):
        raise AssertionError(f"{path.name}: all RuntimeSafe fields must be required")
    for field, schema in (json_schema.get("properties") or {}).items():
        if schema.get("type") not in {"string", "integer", "number", "boolean"}:
            raise AssertionError(f"{path.name}: {field} is not primitive")
    return {
        "kind": kind,
        "path": path,
        "outer": outer,
        "inner": inner,
        "record": {
            "logicalKind": kind,
            "realWidgetId": real_id,
            "widgetName": name,
            "exportSha256": sha256_bytes(raw),
            "viewSha256": sha256_bytes(str(inner.get("view") or "").encode("utf-8")),
            "schemaSha256": sha256_bytes(canonical_json(json_schema)),
            "environment": None,
            "exportedAt": exported_at(path),
            "fieldNames": schema_fields,
            "fieldTypes": {key: value.get("type") for key, value in json_schema["properties"].items()},
            "sourceFile": path.name,
            "sourceType": "tencent-widget-export",
        },
    }


def assert_current_runtime_id(path_name: str, real_id: str, stale_ids: set[str]) -> None:
    if not re.fullmatch(r"[0-9a-f]{32}", str(real_id or "")):
        raise AssertionError(f"{path_name}: real WidgetID is empty or invalid")
    if real_id in stale_ids:
        raise AssertionError(f"{path_name}: stale logical-registry WidgetID is not an environment binding")


def workflow_from_zip(path: Path) -> tuple[dict, str]:
    with zipfile.ZipFile(path) as reader:
        names = [name for name in reader.namelist() if name.endswith("_workflow.json")]
        if len(names) != 1:
            raise AssertionError(f"{path.name}: expected one workflow JSON")
        return json.loads(reader.read(names[0]).decode("utf-8")), names[0]


def discover_workflows(directory: Path) -> dict[str, Path]:
    matches: dict[str, list[Path]] = {key: [] for key in WORKFLOW_NAMES}
    for path in directory.glob("*.zip"):
        try:
            workflow, _ = workflow_from_zip(path)
        except Exception:
            continue
        name = str(workflow.get("WorkflowName") or "")
        for key, pattern in WORKFLOW_SOURCE_PATTERNS.items():
            if pattern.search(name):
                matches[key].append(path)
    selected = {}
    for key, candidates in matches.items():
        if not candidates:
            raise AssertionError(f"real Tencent workflow export missing for {key}")
        candidates.sort(key=lambda item: (item.stat().st_mtime_ns, item.name), reverse=True)
        selected[key] = candidates[0]
    return selected


def schedule_record(workflow_path: Path, logical_registry: dict, environment: str) -> dict:
    workflow, _ = workflow_from_zip(workflow_path)
    widgets = [node for node in workflow.get("Nodes", []) if node.get("NodeType") == "WIDGET"]
    ids = {str((node.get("WidgetNodeData") or {}).get("WidgetID") or "") for node in widgets}
    ids.discard("")
    expected = logical_registry["widgets"]["Schedule"]
    if ids != {expected["widgetId"]}:
        raise AssertionError("01 Schedule export does not reference the verified runtime WidgetID")
    field_sets = {
        tuple(sorted(param.get("ParamName") for param in (node.get("WidgetNodeData") or {}).get("WidgetParam", [])))
        for node in widgets
    }
    if len(field_sets) != 1:
        raise AssertionError("01 Schedule Widget field sets drifted across DAY/WEEK/DATE")
    view_path = ROOT / "contracts" / "schedule" / "view.txt"
    schema_path = ROOT / "contracts" / "schedule" / "schema.json"
    return {
        "logicalKind": "Schedule",
        "realWidgetId": expected["widgetId"],
        "widgetName": expected["name"],
        "exportSha256": expected["sourceSha256"],
        "viewSha256": sha256_file(view_path),
        "schemaSha256": sha256_file(schema_path),
        "environment": environment,
        "exportedAt": exported_at(workflow_path),
        "fieldNames": list(next(iter(field_sets))),
        "fieldTypes": {},
        "sourceFile": workflow_path.name,
        "sourceType": "verified-schedule-export-plus-current-workflow-reference",
        "workflowEvidenceSha256": sha256_file(workflow_path),
    }


def build_runtime_registry(widget_dir: Path, workflows: dict[str, Path], environment: str) -> dict:
    logical = json.loads(LOGICAL_REGISTRY_PATH.read_text(encoding="utf-8"))
    stale_ids = {
        value.get("widgetId") for key, value in logical.get("widgets", {}).items()
        if key != "Schedule" and value.get("widgetId")
    }
    parsed = {}
    for path in widget_dir.glob("*.widget"):
        item = parse_widget(path)
        kind = item["kind"]
        if kind in parsed:
            raise AssertionError(f"duplicate current-environment Widget export: {kind}")
        assert_current_runtime_id(path.name, item["record"]["realWidgetId"], stale_ids)
        item["record"]["environment"] = environment
        parsed[kind] = item
    expected = set(RUNTIME_KINDS) - {"Schedule"}
    if set(parsed) != expected:
        missing = sorted(expected - set(parsed))
        extra = sorted(set(parsed) - expected)
        raise AssertionError(f"runtime Widget export set mismatch; missing={missing}; extra={extra}")
    widgets = {"Schedule": schedule_record(workflows["01"], logical, environment)}
    widgets.update({kind: parsed[kind]["record"] for kind in sorted(parsed)})
    ids = [record["realWidgetId"] for record in widgets.values()]
    if len(ids) != len(set(ids)):
        raise AssertionError("runtime Widget IDs must be unique")
    return {
        "schema": "fosuclass-widget-runtime-registry/v1",
        "logicalRegistry": "competition/adp-kit/widget/native/widget-registry.json",
        "environment": environment,
        "bindingPolicy": "FAIL_CLOSED_CURRENT_ENVIRONMENT_EXPORTS_ONLY",
        "widgets": widgets,
    }


ACTION_PROTOCOL_WRAPPER = r'''

# CampusFlow R4 compiler wrapper: preserve the natural query and append a
# deterministic Action Protocol V2 envelope to every sys.chat message.
_campusflow_runtime_main = main

def main(params: dict) -> dict:
    result = _campusflow_runtime_main(params)
    if not isinstance(result, dict):
        return result
    body = params.get("tool_body") if isinstance(params, dict) else {}
    body = body if isinstance(body, dict) else {}
    query_context = body.get("query") if isinstance(body.get("query"), dict) else {}
    resolved = body.get("resolvedEntity") if isinstance(body.get("resolvedEntity"), dict) else {}
    compared = body.get("compared") if isinstance(body.get("compared"), list) else []

    def infer(label, natural):
        value = f"{label} {natural}"
        if "风险" in value or "冲突" in value or "赶场" in value:
            return "schedule_risk_check"
        if "当天" in value:
            return "schedule_day"
        if "整周" in value:
            return "schedule_week"
        if "课表" in value:
            return "schedule_day"
        if "空教室" in value or "校区" in value or "时段" in value:
            return "classroom_find"
        if "日期" in value or "规划" in value or "计划" in value:
            return "day_plan_choose_date"
        if "选择" in value or "确认" in value:
            return "entity_choose"
        return "recovery_continue"

    def clean(value):
        return str(value if value is not None else "").replace("|", " ").replace("\n", " ").strip()

    entity = resolved
    if not entity and compared and isinstance(compared[0], dict):
        entity = compared[0]
    for key in list(result):
        if not key.endswith("Message") or not isinstance(result.get(key), str):
            continue
        natural = clean(result[key])
        if not natural or "【小序ActionV2】" in natural:
            continue
        label = clean(result.get(key[:-7] + "Label"))
        parts = [f"query={natural}", f"intent={infer(label, natural)}"]
        context = {
            "entityType": entity.get("type") or query_context.get("entityType"),
            "entityName": entity.get("name") or query_context.get("entityName"),
            "week": query_context.get("week"),
            "weekday": query_context.get("weekday"),
            "date": query_context.get("date"),
        }
        for context_key in ("entityType", "entityName", "week", "weekday", "date"):
            value = context.get(context_key)
            if value not in (None, "", 0):
                parts.append(f"{context_key}={clean(value)}")
        result[key] = natural + "\n【小序ActionV2】" + "|".join(parts)
    return result
'''


def node_widget_fields(node: dict) -> list[str]:
    return [param.get("ParamName") for param in (node.get("WidgetNodeData") or {}).get("WidgetParam", [])]


def bind_literal_widget_actions(node: dict) -> None:
    params = (node.get("WidgetNodeData") or {}).get("WidgetParam", [])
    by_name = {param.get("ParamName"): param for param in params}
    for name, param in by_name.items():
        if not str(name or "").endswith("Message"):
            continue
        input_value = param.get("Input") or {}
        if input_value.get("InputType") != "USER_INPUT":
            continue
        user_value = input_value.get("UserInputValue") or {}
        values = user_value.get("Values") or []
        if not values or not isinstance(values[0], str) or not values[0].strip():
            continue
        natural = values[0].replace("|", " ").replace("\n", " ").strip()
        if "【小序ActionV2】" in natural:
            continue
        label_name = name[:-7] + "Label"
        label_values = (((by_name.get(label_name) or {}).get("Input") or {}).get("UserInputValue") or {}).get("Values") or []
        label = str(label_values[0] if label_values else "")
        intent = "entity_choose" if "选择" in label or "确认" in label else "recovery_continue"
        values[0] = f"{natural}\n【小序ActionV2】query={natural}|intent={intent}"


def bind_workflow(key: str, source: Path, registry: dict, environment: str) -> dict:
    workflow, _ = workflow_from_zip(source)
    workflow["WorkflowID"] = str(uuid.uuid5(R4_NAMESPACE, f"{environment}:{key}:R3"))
    workflow["WorkflowName"] = WORKFLOW_NAMES[key]
    workflow["WorkflowDesc"] = (
        f"CampusFlow ADP R4 environment-bound {key}; current Tencent export base; "
        "RuntimeSafe primitive Widget boundary; Action Protocol V2."
    )
    workflow["ReleaseTime"] = ""
    runtime_by_fields = {
        tuple(sorted(record["fieldNames"])): record
        for record in registry["widgets"].values()
    }
    widget_count = 0
    for node in workflow.get("Nodes", []):
        if node.get("NodeType") == "WIDGET":
            widget_count += 1
            fields = tuple(sorted(node_widget_fields(node)))
            record = runtime_by_fields.get(fields)
            if not record:
                raise AssertionError(f"{key}/{node.get('NodeName')}: WidgetParam does not match runtime registry")
            data = node.setdefault("WidgetNodeData", {})
            data["WidgetID"] = record["realWidgetId"]
            if not data["WidgetID"]:
                raise AssertionError(f"{key}/{node.get('NodeName')}: empty WidgetID")
            for param in data.get("WidgetParam", []):
                if param.get("ParamType") not in {"STRING", "INT", "FLOAT", "BOOL"}:
                    raise AssertionError(f"{key}/{node.get('NodeName')}: non-primitive Widget field")
            if key in {"02", "03", "04"}:
                bind_literal_widget_actions(node)
        code_data = node.get("CodeExecutorNodeData") or {}
        code = code_data.get("Code")
        if key in {"02", "03", "04"} and isinstance(code, str) and "Message\"" in code and "_campusflow_runtime_main" not in code:
            code_data["Code"] = code.rstrip() + ACTION_PROTOCOL_WRAPPER
    if not widget_count:
        raise AssertionError(f"{key}: no WIDGET nodes")
    allowed_ids = {record["realWidgetId"] for record in registry["widgets"].values()}
    for node in workflow.get("Nodes", []):
        if node.get("NodeType") != "WIDGET":
            continue
        widget_id = (node.get("WidgetNodeData") or {}).get("WidgetID")
        if not widget_id or widget_id not in allowed_ids:
            raise AssertionError(f"{key}/{node.get('NodeName')}: WidgetID not in runtime registry")
    return workflow


def workbook_rows(source: Path, name: str) -> tuple[list, list[list]]:
    with zipfile.ZipFile(source) as reader:
        workbook = load_workbook(io.BytesIO(reader.read(name)), read_only=True, data_only=False)
    rows = list(workbook.active.iter_rows(values_only=True))
    if not rows:
        raise AssertionError(f"{source.name}/{name}: empty workbook")
    return list(rows[0]), [list(row) for row in rows[1:]]


def deterministic_xlsx(headers: list, rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    fixed = datetime.datetime(2026, 8, 14, 0, 0, 0)
    workbook.properties.creator = "FosuClass ADP R4 Binder"
    workbook.properties.lastModifiedBy = "FosuClass ADP R4 Binder"
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    with tempfile.TemporaryDirectory() as temp_dir:
        raw_path = Path(temp_dir) / "raw.xlsx"
        output_path = Path(temp_dir) / "deterministic.xlsx"
        workbook.save(raw_path)
        with zipfile.ZipFile(raw_path) as reader:
            entries = [(name, reader.read(name)) for name in sorted(reader.namelist())]
        with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as writer:
            for name, data in entries:
                if name == "docProps/core.xml":
                    data = re.sub(
                        rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                        rb"\g<1>2026-08-14T00:00:00Z\g<2>",
                        data,
                    )
                info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o644 << 16
                writer.writestr(info, data)
        return output_path.read_bytes()


def bound_workbooks(source: Path, workflow: dict, key: str) -> dict[str, bytes]:
    workflow_id = workflow["WorkflowID"]
    payloads = {}
    expected_headers = {
        "workflows.xlsx": ["WorkflowId", "WorkflowName", "WorkflowDescription", "CanvasStructure"],
        "parameters.xlsx": ["WorkflowId", "WorkflowNodeId", "WorkflowNodeName", "ParameterId", "ParameterName", "ParameterDescription", "ParameterType", "ParameterCorrectExample", "ParameterIncorrectExample", "ParameterParentId", "ParameterRequiredStatus"],
        "example_queries.xlsx": ["WorkflowId", "ExampleQueryContent"],
        "variables.xlsx": ["WorkflowId", "VariableId", "VariableName", "VariableDescription", "VariableType", "VariableDefaultValue", "VariableDefaultValueFileName", "ParameterType"],
        "workflow_references.xlsx": ["WorkflowId", "WorkflowNodeId", "ReferenceWorkflowId", "ReferenceWorkflowName"],
    }
    for name in ROOT_FILES:
        # Recent Tencent re-exports may collapse every workbook to the single
        # WorkflowId header.  The real import contract is the full header set
        # captured by the working 01 package, so never propagate that lossy
        # export representation into a generated import artifact.
        source_headers, source_rows = workbook_rows(source, name)
        headers = expected_headers[name]
        rows = source_rows if source_headers == headers else []
        if name == "workflows.xlsx":
            rows = [[workflow_id, workflow["WorkflowName"], workflow["WorkflowDesc"], f"{workflow_id}_workflow.json"]]
        elif name == "example_queries.xlsx":
            rows = [[workflow_id, example] for example in WORKFLOW_EXAMPLES[key]]
        elif name == "variables.xlsx" and key == "01" and not rows:
            rows = [[
                workflow_id,
                "a39583c8-2484-46f2-8ba8-e7384a0365e8",
                "campus_api_authorization",
                "导入后绑定当前赛事空间环境变量；占位值不是凭据",
                "STRING",
                "Bearer ADP_IMPORT_PLACEHOLDER_NOT_A_CREDENTIAL",
                "",
                "1",
            ]]
        else:
            if rows:
                workflow_index = headers.index("WorkflowId")
                for row in rows:
                    row[workflow_index] = workflow_id
        payloads[name] = deterministic_xlsx(headers, rows)
    return payloads


def write_workflow_zip(path: Path, workflow: dict, workbooks: dict[str, bytes]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workflow_name = f"{workflow['WorkflowID']}_workflow.json"
    payloads = {
        workflow_name: json.dumps(workflow, ensure_ascii=False, separators=(",", ":")).encode("utf-8"),
        **workbooks,
    }
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as writer:
        for name in [workflow_name, *ROOT_FILES]:
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            writer.writestr(info, payloads[name])


def bind_environment(widget_dir: Path, workflow_dir: Path, environment: str, output_dir: Path) -> dict:
    if not widget_dir.is_dir() or not workflow_dir.is_dir():
        raise AssertionError("widget-dir and workflow-dir must exist")
    workflows = discover_workflows(workflow_dir)
    registry = build_runtime_registry(widget_dir, workflows, environment)
    output_dir.mkdir(parents=True, exist_ok=True)
    registry_path = output_dir / "widget-registry.runtime.json"
    registry_path.write_text(json.dumps(registry, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    artifacts = []
    for key in WORKFLOW_NAMES:
        workflow = bind_workflow(key, workflows[key], registry, environment)
        target = output_dir / WORKFLOW_FILES[key]
        write_workflow_zip(target, workflow, bound_workbooks(workflows[key], workflow, key))
        artifacts.append({
            "slot": key,
            "file": target.name,
            "workflowId": workflow["WorkflowID"],
            "workflowName": workflow["WorkflowName"],
            "sourceFile": workflows[key].name,
            "sourceSha256": sha256_file(workflows[key]),
            "sha256": sha256_file(target),
        })
    return {"registry": registry, "registryPath": registry_path, "artifacts": artifacts}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--widget-dir", required=True)
    parser.add_argument("--workflow-dir", required=True)
    parser.add_argument("--env", required=True)
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT))
    args = parser.parse_args()
    result = bind_environment(
        Path(args.widget_dir).resolve(),
        Path(args.workflow_dir).resolve(),
        args.env,
        Path(args.output_dir).resolve(),
    )
    print(f"ADP runtime binding: PASS ({args.env})")
    for item in result["artifacts"]:
        print(f"{item['slot']}: {item['workflowName']} -> {item['file']} [{item['sha256']}]")


if __name__ == "__main__":
    main()
