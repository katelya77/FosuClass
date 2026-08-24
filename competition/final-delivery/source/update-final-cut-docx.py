from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "校园智序-小序-智能体设计说明书.docx"
WIDGET_IMAGE = ROOT / "assets" / "doc-p15-real-widget.png"
EXECUTION_IMAGE = ROOT / "assets" / "doc-p18-execution-evidence.png"


def blocks(document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def first_heading_index(items, prefix):
    for index, item in enumerate(items):
        if isinstance(item, Paragraph) and item.text.strip().startswith(prefix):
            return index
    raise RuntimeError(f"Missing page heading: {prefix}")


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag.endswith("}pPr"):
            continue
        paragraph._p.remove(child)


doc = Document(DOCX)
items = list(blocks(doc))

# Page 15: replace the conceptual slide miniature with the real, anonymized
# native ADP conversation and official Widget render.
p15 = first_heading_index(items, "15  /")
widget_paragraph = next(
    item
    for item in items[p15 + 1 : p15 + 12]
    if isinstance(item, Paragraph) and item._p.xpath(".//a:blip")
)
clear_paragraph(widget_paragraph)
widget_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
widget_paragraph.paragraph_format.space_before = Pt(4)
widget_paragraph.paragraph_format.space_after = Pt(5)
widget_paragraph.add_run().add_picture(str(WIDGET_IMAGE), width=Inches(7.22))

# Page 18: replace the former three-label concept table with a compact strip
# captured from the real Main → Child → Main → CampusTools → Widget evidence slide.
items = list(blocks(doc))
p18 = first_heading_index(items, "18  /")
evidence_paragraph = next(
    (
        item
        for item in items[p18 + 1 : p18 + 10]
        if isinstance(item, Paragraph) and item._p.xpath(".//a:blip")
    ),
    None,
)
if evidence_paragraph is None:
    evidence_table = next(item for item in items[p18 + 1 : p18 + 10] if isinstance(item, Table))
    new_p = OxmlElement("w:p")
    evidence_table._tbl.addprevious(new_p)
    evidence_paragraph = Paragraph(new_p, doc)
    evidence_table._tbl.getparent().remove(evidence_table._tbl)
else:
    clear_paragraph(evidence_paragraph)
evidence_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
evidence_paragraph.paragraph_format.space_before = Pt(3)
evidence_paragraph.paragraph_format.space_after = Pt(5)
evidence_paragraph.add_run().add_picture(str(EXECUTION_IMAGE), width=Inches(7.22))
caption_run = evidence_paragraph.add_run("\n真实运行证据：Main → Child → Main；领域 Agent 调用 CampusTools，最终由官方 Widget 投影 Verified 结果。")
caption_run.font.name = "Microsoft YaHei"
caption_run.font.size = Pt(8.5)
caption_run.font.color.rgb = RGBColor(127, 111, 105)

doc.core_properties.title = "校园智序·小序—智能体设计说明书"
doc.core_properties.subject = "比赛最终交付 · Evidence-first Final Cut"
doc.core_properties.author = "校园智序·小序"
doc.core_properties.last_modified_by = "校园智序·小序"
doc.core_properties.comments = ""

temp = DOCX.with_name("final-cut-design-guide.docx")
doc.save(temp)
temp.replace(DOCX)
print(f"updated={DOCX}")
